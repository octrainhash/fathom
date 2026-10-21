# Copyright (C) 2026 The Fathomist. Licensed under AGPL-3.0-or-later (see LICENSE).
"""Everything the assistant can do. Each tool is a plain function plus a description the model reads.

risk levels (decide when the user is asked first):
  read    only looks (read files, search, web)             never asks (except in 'ask about everything' mode)
  create  makes NEW files in the assistant's output folder never asks unless 'ask about everything'
  write   changes, moves, overwrites or deletes your files, or runs a Fathom tool that does   asks unless 'never ask'
  run     runs code or commands, adds or changes a tool    asks unless 'never ask'"""
from __future__ import annotations

import base64
import difflib
import html as _html
import io
import json
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import typing
import urllib.parse
from dataclasses import dataclass, field
from pathlib import Path

from pydantic import BaseModel

from app.ai import catalog, docs, media, plugins, store
from app.runtime import no_window_flags
from app.service import UserError

MAX_TEXT = 60000


@dataclass
class Result:
    text: str = ""
    images: list = field(default_factory=list)       # [(media_type, base64)]
    files: list = field(default_factory=list)        # [{"path","name","kind"}]
    proof: list = field(default_factory=list)        # proof items for the chat gallery
    ok: bool = True
    ui: dict | None = None                           # an instruction for the window (open a tool, add files)


TOOLS: dict[str, dict] = {}


def tool(name: str, desc: str, props: dict, required: list[str] | None = None, risk: str = "read"):
    def deco(fn):
        TOOLS[name] = {"name": name, "description": desc, "input_schema": {"type": "object", "properties": props, "required": required or []}, "risk": risk, "fn": fn}
        return fn
    return deco


def schemas() -> list[dict]:
    return [{"name": t["name"], "description": t["description"], "input_schema": t["input_schema"]} for t in TOOLS.values()]


# Small local models (Ollama) have a short memory and read the prompt slowly, so they get fewer tools with shorter descriptions.
LOCAL_TOOLS = ("get_screen_context", "list_files", "read_file", "view_image", "search_files", "write_file", "edit_file", "file_operation",
               "create_document", "create_presentation", "create_spreadsheet", "fathom_tools", "fathom_tool_info", "fathom_run",
               "find_in_fathom", "open_in_fathom", "web_search", "web_fetch", "run_python", "ask_user", "show_files", "remember")


def _short(text: str, n: int) -> str:
    text = " ".join(str(text or "").split())
    if len(text) <= n:
        return text
    cut = text[:n]
    stop = max(cut.rfind(". "), cut.rfind("; "))
    return (cut[:stop + 1] if stop > n // 2 else cut.rstrip(" ,;:") + "...")


def _slim_schema(node):
    if isinstance(node, dict):
        out = {}
        for k, v in node.items():
            if k == "description" and isinstance(v, str):
                out[k] = _short(v, 90)
            else:
                out[k] = _slim_schema(v)
        return out
    if isinstance(node, list):
        return [_slim_schema(x) for x in node]
    return node


def schemas_local() -> list[dict]:
    """The same tools, but only the ones a small local model needs, with short descriptions."""
    out = []
    for name in LOCAL_TOOLS:
        t = TOOLS.get(name)
        if t:
            out.append({"name": name, "description": _short(t["description"], 170), "input_schema": _slim_schema(t["input_schema"])})
    return out


def risk_of(name: str, inp: dict) -> str:
    t = TOOLS.get(name)
    if not t:
        return "read"
    if name == "fathom_run":
        return _route_risk(str(inp.get("path", "")), str(inp.get("method", "POST")))
    if name == "write_file" and not Path(os.path.expanduser(str(inp.get("path", "")))).exists():
        p = _resolve(str(inp.get("path", "")))
        return "create" if _inside_output(p) else "write"
    return t["risk"]


def _inside_output(p: Path) -> bool:
    try:
        p.resolve().relative_to(store.output_dir().resolve())
        return True
    except Exception:
        return False


READ_ROUTE = re.compile(r"/(preview|info|hash|meta$|kinds|plan|search|undo$|status|files$|items|printers|drives|health|storage$|formats|capabilities|compat$|advise|space|hex/view|hex/find|qr/read|compare|scan$|exif$|perms$|times$|rename/preview|organize/preview|emptyfolders$|ads/scan|attrs/find|carve/kinds|damage/modes|torrent/info|dups$)")


def _route_risk(path: str, method: str) -> str:
    if method.upper() == "GET":
        return "read"
    if path.startswith("/api/tools/scan") or path.startswith("/api/tools/exif") or path.startswith("/api/tools/perms") or path.startswith("/api/tools/times"):
        return "write"
    return "read" if READ_ROUTE.search(path) else "write"


# ------------------------------------------------------------------ helpers
def _resolve(p: str, sess=None) -> Path:
    p = (p or "").strip().strip('"')
    if not p:
        raise UserError("No path was given.")
    q = Path(os.path.expandvars(os.path.expanduser(p)))
    if not q.is_absolute():
        base = store.output_dir()
        q = base / q
    return q


def _clip(s: str, n: int = MAX_TEXT) -> str:
    return s if len(s) <= n else s[:n] + f"\n… [{len(s) - n} more characters; ask for a smaller part, e.g. with pages or offset]"


def _kind(p: Path) -> str:
    e = p.suffix.lower()
    for k, ext in plugins.KINDS.items():
        if e in ext:
            return k
    return "file"


def _file_rec(p: Path) -> dict:
    return {"path": str(p), "name": p.name, "kind": _kind(p), "bytes": p.stat().st_size if p.exists() and p.is_file() else 0}


def _new_file(name: str, ext: str, sess=None) -> Path:
    name = re.sub(r'[\\/:*?"<>|\r\n]+', " ", name or "").strip(" .") or "Untitled"
    if not name.lower().endswith(ext.lower()):
        name = re.sub(r"\.\w{1,5}$", "", name) + ext
    d = store.output_dir()
    p = d / name
    stem, i = p.stem, 2
    while p.exists():
        p = d / f"{stem} ({i}){ext}"
        i += 1
    return p


def _backup(p: Path) -> str:
    if p.exists() and p.is_file():
        b = store.output_dir() / "backups"
        b.mkdir(exist_ok=True)
        dst = b / f"{time.strftime('%Y%m%d-%H%M%S')} {p.name}"
        shutil.copy2(p, dst)
        return str(dst)
    return ""


def _image_block(p: Path, page: int = 1, max_side: int = 1568) -> tuple[str, str]:
    from PIL import Image

    if p.suffix.lower() == ".pdf" or p.suffix.lower() in plugins.KINDS["office"]:
        from app import pdfview
        data, _m = pdfview.page_png(str(p), max(0, page - 1), max_side)
        im = Image.open(io.BytesIO(data))
    else:
        im = Image.open(p)
        try:
            im.seek(max(0, page - 1))
        except Exception:
            pass
    im = im.convert("RGBA") if im.mode in ("P", "LA", "RGBA") else im.convert("RGB")
    if im.mode == "RGBA":
        bg = Image.new("RGB", im.size, "white")
        bg.paste(im, mask=im.getchannel("A"))
        im = bg
    if max(im.size) > max_side:
        im.thumbnail((max_side, max_side), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=88)
    return "image/jpeg", base64.b64encode(buf.getvalue()).decode()


# ================================================================== looking around
@tool("get_screen_context", "What the user is doing in Fathom right now: the open tab, the values in its form, the files added on the left, the folder results go to, recent actions and results. Call this first whenever the request is about 'this', 'these files', 'what I'm doing' or you are unsure what they mean.", {})
def get_screen_context(inp, sess):
    from app.ai import agent
    return Result(json.dumps(agent.context_for(sess), ensure_ascii=False, indent=1)[:MAX_TEXT])


@tool("list_files", "List files and folders in a folder (name, size, date). Use recursive to go deeper and pattern for a name filter like *.pdf.",
      {"path": {"type": "string"}, "pattern": {"type": "string", "description": "glob such as *.pdf"}, "recursive": {"type": "boolean"}, "limit": {"type": "integer"}}, ["path"])
def list_files(inp, sess):
    p = _resolve(inp["path"])
    if p.is_file():
        st = p.stat()
        return Result(f"{p} is a file, {st.st_size} bytes.")
    if not p.exists():
        raise UserError(f"Folder not found: {p}")
    pat = inp.get("pattern") or "*"
    limit = int(inp.get("limit") or 300)
    it = p.rglob(pat) if inp.get("recursive") else p.glob(pat)
    rows = []
    total = 0
    for f in it:
        total += 1
        if len(rows) < limit:
            try:
                st = f.stat()
                rows.append(f"{'[dir] ' if f.is_dir() else ''}{f.relative_to(p)}\t{'' if f.is_dir() else st.st_size}\t{time.strftime('%Y-%m-%d %H:%M', time.localtime(st.st_mtime))}")
            except OSError:
                continue
    return Result(f"{total} item(s) in {p}" + (f" (showing {limit})" if total > limit else "") + "\n" + "\n".join(rows))


def _pdf_text(p: Path, pages: str, ocr: bool) -> str:
    from pypdf import PdfReader

    from app import pagesel
    r = PdfReader(str(p))
    if r.is_encrypted:
        try:
            r.decrypt("")
        except Exception:
            raise UserError("This PDF has a password. Ask the user for it, then use the PDF password tool.")
    n = len(r.pages)
    sel = pagesel.parse(pages, n) if pages else list(range(1, n + 1))
    out = []
    for pg in sel[:400]:
        t = (r.pages[pg - 1].extract_text() or "").strip()
        if len(t) < 25 and ocr:
            t = _ocr_page(p, pg) or t
        out.append(f"--- page {pg} of {n} ---\n{t}")
    return "\n".join(out)


def _ocr_page(p: Path, pg: int) -> str:
    try:
        from PIL import Image

        from app import pdfview
        from app.ocr import ocr_boxes
        data, _m = pdfview.page_png(str(p), pg - 1, 1800)
        boxes = ocr_boxes(Image.open(io.BytesIO(data)).convert("RGB"))
        boxes = sorted(boxes, key=lambda b: (round(b.get("y", 0) / 14), b.get("x", 0)))
        return " ".join(b.get("text", "") for b in boxes)
    except Exception:
        return ""


def read_any(p: Path, pages: str = "", ocr: bool = True) -> str:
    e = p.suffix.lower()
    if e == ".pdf":
        return _pdf_text(p, pages, ocr)
    if e == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [x.text for x in d.paragraphs]
        for ti, t in enumerate(d.tables, 1):
            parts.append(f"[table {ti}]")
            parts += [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        return "\n".join(parts)
    if e == ".pptx":
        from pptx import Presentation
        out = []
        for i, s in enumerate(Presentation(str(p)).slides, 1):
            out.append(f"--- slide {i} ---")
            for sh in s.shapes:
                if sh.has_text_frame:
                    out.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    out += [" | ".join(c.text for c in r.cells) for r in sh.table.rows]
            if s.has_notes_slide:
                out.append("[notes] " + s.notes_slide.notes_text_frame.text)
        return "\n".join(out)
    if e in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(str(p), data_only=False, read_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"--- sheet {ws.title} ({ws.max_row} rows x {ws.max_column} cols) ---")
            for r in ws.iter_rows(min_row=1, max_row=300, values_only=True):
                out.append("\t".join("" if v is None else str(v) for v in r))
        return "\n".join(out)
    if e in (".html", ".htm"):
        from bs4 import BeautifulSoup
        s = BeautifulSoup(p.read_text("utf-8", "replace"), "html.parser")
        for x in s(["script", "style"]):
            x.decompose()
        return re.sub(r"\n{3,}", "\n\n", s.get_text("\n"))
    if e in (".doc", ".xls", ".ppt", ".odt", ".ods", ".odp", ".rtf"):
        from app import pdfview
        return _pdf_text(pdfview._as_pdf(p), pages, ocr)
    if e in (".zip",):
        import zipfile
        with zipfile.ZipFile(p) as z:
            return "\n".join(f"{i.file_size}\t{i.filename}" for i in z.infolist()[:500])
    if e in plugins.KINDS["image"]:
        t = _ocr_whole_image(p) if ocr else ""
        return "(This is a picture. Use view_image to look at it.)\n" + (f"Text found in it:\n{t}" if t else "")
    raw = p.read_bytes()[:20_000_000]
    if b"\x00" in raw[:4096]:
        return f"(Binary file, {p.stat().st_size} bytes. Use run_python or a Fathom tool such as hex view to inspect it.)"
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(enc)
        except Exception:
            continue
    return raw.decode("utf-8", "replace")


def _ocr_whole_image(p: Path) -> str:
    try:
        from PIL import Image

        from app.ocr import ocr_boxes
        boxes = ocr_boxes(Image.open(p).convert("RGB"))
        boxes = sorted(boxes, key=lambda b: (round(b.get("y", 0) / 14), b.get("x", 0)))
        return " ".join(b.get("text", "") for b in boxes)
    except Exception:
        return ""


@tool("read_file", "Read a file as text: PDF (with OCR for scanned pages), Word, PowerPoint, Excel (formulas shown), CSV, JSON, web pages, code, plain text, zip listings. Use pages ('1-3') for PDFs and offset for long text.",
      {"path": {"type": "string"}, "pages": {"type": "string"}, "offset": {"type": "integer"}, "max_chars": {"type": "integer"}}, ["path"])
def read_file(inp, sess):
    p = _resolve(inp["path"])
    if not p.exists():
        raise UserError(f"File not found: {p}")
    if p.is_dir():
        return list_files({"path": str(p)}, sess)
    t = read_any(p, str(inp.get("pages") or ""))
    off = int(inp.get("offset") or 0)
    n = min(int(inp.get("max_chars") or MAX_TEXT), 200000)
    chunk = t[off:off + n]
    more = len(t) - off - len(chunk)
    return Result(chunk + (f"\n… [{more} more characters; call again with offset={off + len(chunk)}]" if more > 0 else ""))


@tool("view_image", "Look at a picture, or one page of a PDF / Word / PowerPoint file, as an image (you can see it). Use it to check how a document or design looks, to read handwriting, to understand a photo or screenshot.",
      {"path": {"type": "string"}, "page": {"type": "integer", "description": "page number, 1-based"}, "detail": {"type": "string", "enum": ["normal", "high"]}}, ["path"])
def view_image(inp, sess):
    p = _resolve(inp["path"])
    if not p.exists():
        raise UserError(f"File not found: {p}")
    mt, b64 = _image_block(p, int(inp.get("page") or 1), 2000 if inp.get("detail") == "high" else 1568)
    return Result(f"Here is {p.name}" + (f", page {inp.get('page')}" if inp.get("page") else "") + ".", images=[(mt, b64)])


@tool("look_at_screen", "See the user's screen right now (all monitors) as a picture. Only works while the user has pressed 'Watch my screen'. Use it to see what they are doing in Fathom or any other program, to follow a job's progress, or to answer 'what is this?' / 'what do I click?'. Look again whenever the screen may have changed.", {})
def look_at_screen(inp, sess):
    from app.ai import screen
    if not screen.is_on():
        raise UserError("Screen watching is off. Ask the user to press the 'Watch my screen' button in the chat if they want you to see their screen.")
    f = screen.grab()
    return Result(f"This is the user's screen at {time.strftime('%H:%M:%S')} ({f['width']}x{f['height']} px).", images=[(f["media_type"], f["data"])])


@tool("search_files", "Use Fathom's own search: find words inside files (PDF, Word, text, pictures with OCR, ...) or objects in pictures and videos. mode 'text' or 'object'. Returns the matches with file and place.",
      {"query": {"type": "string"}, "paths": {"type": "array", "items": {"type": "string"}}, "mode": {"type": "string", "enum": ["text", "object"]}, "options": {"type": "object", "description": "extra search options such as ocr, case_sensitive, use_regex, whole_word"}}, ["query", "paths"])
def search_files(inp, sess):
    from app import service
    opts = dict(inp.get("options") or {})
    opts.setdefault("mode", inp.get("mode") or "text")
    r = service.create_search(inp["query"], [str(_resolve(x)) for x in inp["paths"]], opts)
    job = _wait_job(sess, r["job"], search=True, timeout=900)
    return Result(job)


def _wait_job(sess, job_id: str, search: bool = False, timeout: float = 3600) -> str:
    from app import ops, service
    t0 = time.time()
    since, results = 0, []
    snap = {}
    while True:
        if sess is not None and sess.cancel.is_set():
            (service.cancel_job if search else ops.cancel)(job_id)
            return "Stopped by the user."
        snap = service.job_status(job_id, since) if search else ops.OPJOBS[job_id].snapshot(since)
        results += snap.get("results", [])
        since = snap.get("next", since)
        if sess is not None:
            sess.progress(snap.get("message") or snap.get("current") or "", snap.get("progress"))
        if snap.get("status") in ("done", "cancelled", "failed"):
            break
        if time.time() - t0 > timeout:
            return f"Still running after {int(timeout)} s (job {job_id}); check it later with fathom_run GET /api/{'jobs' if search else 'ops'}/{job_id}."
        time.sleep(0.4)
    return _summarise_job(snap, results, search)


def _strip_big(o, depth=0):
    if isinstance(o, dict):
        return {k: ("[picture data omitted]" if isinstance(v, str) and v.startswith("data:image") else _strip_big(v, depth + 1)) for k, v in o.items() if k not in ("seq",)}
    if isinstance(o, list):
        return [_strip_big(x, depth + 1) for x in o[:60]]
    if isinstance(o, str) and len(o) > 600:
        return o[:600] + "…"
    return o


def _summarise_job(snap: dict, results: list, search: bool) -> str:
    head = f"Job {snap.get('status')}: {snap.get('message') or ''}".strip()
    if search:
        lines = [head, f"{len(results)} match(es)"]
        for r in results[:80]:
            lines.append(json.dumps(_strip_big(r), ensure_ascii=False)[:420])
        if snap.get("error_count"):
            lines.append(f"{snap['error_count']} file(s) could not be read.")
        return _clip("\n".join(lines), 30000)
    counts = snap.get("counts") or {}
    lines = [head, f"done {counts.get('done', 0)}, errors {counts.get('error', 0)}, skipped {counts.get('skipped', 0)}"]
    for r in results[:60]:
        outs = ", ".join(o["path"] for o in r.get("outputs", [])[:4])
        extra = {k: v for k, v in (r.get("extra") or {}).items() if k != "proof"}
        lines.append(f"- {r['name']}: {r['status']}" + (f" → {outs}" if outs else "") + (f" | {r['note']}" if r.get("note") else "") + (f" | ERROR: {r['error']}" if r.get("error") else "") + (f" | {json.dumps(_strip_big(extra), ensure_ascii=False)[:300]}" if extra else ""))
    if len(results) > 60:
        lines.append(f"… and {len(results) - 60} more results")
    if snap.get("data"):
        lines.append("data: " + json.dumps(_strip_big(snap["data"]), ensure_ascii=False)[:1500])
    return _clip("\n".join(lines), 30000)


# ================================================================== changing files
@tool("write_file", "Write a text file (new, or overwrite; the old version is saved to a backups folder first). Relative paths go in the assistant's output folder.", {"path": {"type": "string"}, "content": {"type": "string"}, "append": {"type": "boolean"}}, ["path", "content"], risk="write")
def write_file(inp, sess):
    p = _resolve(inp["path"])
    p.parent.mkdir(parents=True, exist_ok=True)
    b = _backup(p) if p.exists() else ""
    if inp.get("append") and p.exists():
        with open(p, "a", encoding="utf-8", newline="") as f:
            f.write(inp["content"])
    else:
        p.write_text(inp["content"], encoding="utf-8", newline="")
    return Result(f"Wrote {p} ({p.stat().st_size} bytes)." + (f" Previous version saved at {b}." if b else ""), files=[_file_rec(p)])


@tool("edit_file", "Change part of a text file: replace exact text old with new (the old version is saved to a backups folder first).", {"path": {"type": "string"}, "old": {"type": "string"}, "new": {"type": "string"}, "all": {"type": "boolean"}}, ["path", "old", "new"], risk="write")
def edit_file(inp, sess):
    p = _resolve(inp["path"])
    t = p.read_text("utf-8")
    n = t.count(inp["old"])
    if not n:
        raise UserError("That exact text was not found in the file.")
    if n > 1 and not inp.get("all"):
        raise UserError(f"That text appears {n} times; give more surrounding text or set all=true.")
    b = _backup(p)
    p.write_text(t.replace(inp["old"], inp["new"]) if inp.get("all") else t.replace(inp["old"], inp["new"], 1), encoding="utf-8", newline="")
    return Result(f"Changed {n if inp.get('all') else 1} place(s) in {p}. Backup: {b}", files=[_file_rec(p)])


@tool("file_operation", "Make a folder, copy, move, rename or delete (to the Recycle Bin) files or folders.", {"action": {"type": "string", "enum": ["make_folder", "copy", "move", "rename", "delete"]}, "path": {"type": "string"}, "to": {"type": "string"}}, ["action", "path"], risk="write")
def file_operation(inp, sess):
    a, p = inp["action"], _resolve(inp["path"])
    if a == "make_folder":
        p.mkdir(parents=True, exist_ok=True)
        return Result(f"Folder ready: {p}")
    if not p.exists():
        raise UserError(f"Not found: {p}")
    if a == "delete":
        from app.recycle import to_recycle_bin
        to_recycle_bin(p)
        return Result(f"Moved to the Recycle Bin: {p}")
    if not inp.get("to"):
        raise UserError("Say where it should go (to).")
    dst = _resolve(inp["to"])
    if a == "rename" and not any(c in inp["to"] for c in "\\/"):
        dst = p.with_name(inp["to"])
    if dst.exists() and dst.is_dir():
        dst = dst / p.name
    if dst.exists():
        raise UserError(f"{dst} already exists; choose another name.")
    dst.parent.mkdir(parents=True, exist_ok=True)
    (shutil.copytree if a == "copy" and p.is_dir() else shutil.copy2 if a == "copy" else shutil.move)(str(p), str(dst))
    return Result(f"{a} done: {p} → {dst}", files=[_file_rec(dst)] if dst.is_file() else [])


# ================================================================== making documents
def _text_of(inp, key="content") -> str:
    s = inp.get(key)
    if s is None:
        raise UserError(f"'{key}' is missing.")
    return str(s)


@tool("create_document", "Make a finished document from Markdown-style text: pdf, docx (Word), html (web page), md or txt. Supports headings, bold/italic, lists (nested, checkboxes), tables, code blocks, quotes, pictures (![caption](path)), page breaks (a line with \\pagebreak), cover page, table of contents, header/footer, themes. "
      "Write the content fully and carefully; it becomes the real file. After making a PDF, look at it with view_image to check the layout.",
      {"format": {"type": "string", "enum": ["pdf", "docx", "html", "md", "txt"]}, "title": {"type": "string"}, "content": {"type": "string", "description": "the whole document in Markdown"}, "filename": {"type": "string"},
       "options": {"type": "object", "description": "theme: clean|report|modern|warm|dark-accent; toc, cover (bool); subtitle, author, date, header_text, footer_text; page: A4|Letter; orientation: portrait|landscape; margin_mm; align: left|justify; size (font pt); accent (hex); base_dir (for relative picture paths)"}}, ["format", "content"], risk="create")
def create_document(inp, sess):
    fmt = inp["format"].lower()
    title = inp.get("title") or ""
    p = _new_file(inp.get("filename") or title or "Document", "." + fmt)
    md = _text_of(inp)
    opts = dict(inp.get("options") or {})
    if fmt == "pdf":
        r = docs.make_pdf(md, p, title, opts)
        extra = f" {r['pages']} page(s)."
    elif fmt == "docx":
        docs.make_docx(md, p, title, opts)
        extra = ""
    elif fmt == "html":
        docs.make_html(md, p, title, opts)
        extra = ""
    else:
        p.write_text(md, encoding="utf-8")
        extra = ""
    return Result(f"Created {p}.{extra}", files=[_file_rec(p)])


@tool("create_presentation", "Make a PowerPoint (.pptx) deck. slides = list of {layout: title|section|bullets|two_column|image|table|chart|quote, title, subtitle, bullets: [text | {text, level}], bullets2, image (path), caption, table {header, rows}, chart {type: bar|barh|line|pie|area, categories, series:[{name, values}]}, quote, author, notes}.",
      {"title": {"type": "string"}, "slides": {"type": "array", "items": {"type": "object"}}, "filename": {"type": "string"}, "options": {"type": "object", "description": "theme, accent, dark (bool), subtitle, base_dir"}}, ["slides"], risk="create")
def create_presentation(inp, sess):
    p = _new_file(inp.get("filename") or inp.get("title") or "Presentation", ".pptx")
    r = docs.make_pptx(inp["slides"], p, inp.get("title") or "", dict(inp.get("options") or {}))
    return Result(f"Created {p} with {r['slides']} slides.", files=[_file_rec(p)])


@tool("create_spreadsheet", "Make an Excel (.xlsx) workbook. sheets = list of {name, rows: [[header...],[row]...] (strings starting with = are formulas), formats: {column letter: number format}, chart: {type: bar|line|pie|scatter, title}}.",
      {"sheets": {"type": "array", "items": {"type": "object"}}, "filename": {"type": "string"}, "options": {"type": "object"}}, ["sheets"], risk="create")
def create_spreadsheet(inp, sess):
    p = _new_file(inp.get("filename") or "Spreadsheet", ".xlsx")
    r = docs.make_xlsx(inp["sheets"], p, dict(inp.get("options") or {}))
    return Result(f"Created {p} with sheets {', '.join(r['sheets'])}.", files=[_file_rec(p)])


@tool("create_chart", "Draw a chart picture (png). spec = {type: bar|barh|line|area|pie|scatter|hist, categories, series:[{name, values}], title, xlabel, ylabel, colors}.", {"spec": {"type": "object"}, "filename": {"type": "string"}}, ["spec"], risk="create")
def create_chart(inp, sess):
    p = _new_file(inp.get("filename") or inp["spec"].get("title") or "Chart", ".png")
    docs.make_chart(inp["spec"], p)
    mt, b64 = _image_block(p)
    return Result(f"Created {p}.", images=[(mt, b64)], files=[_file_rec(p)])


# ================================================================== pictures, voice, video
@tool("design_image", "Draw a poster, card, cover, banner, thumbnail, infographic or diagram from layers (works offline, exact text, any size). spec = {width, height, background: '#hex' | {gradient:[c1,c2],angle} | {image:path,fit:'cover'}, layers:[{type:'text',text,x,y,size,color,bold,font:'sans|serif|mono',align,max_width,shadow} | {type:'rect',x,y,w,h,color,radius,outline,width} | {type:'ellipse',...} | {type:'line',x1,y1,x2,y2,color,width} | {type:'image',path,x,y,w,h,fit,radius,opacity,shadow} | {type:'gradient',...}]}. x:'center' centres. Colours may carry alpha (#rrggbbaa). Look at the result with view_image and fix it if needed.",
      {"spec": {"type": "object"}, "filename": {"type": "string"}, "format": {"type": "string", "enum": ["png", "jpg"]}}, ["spec"], risk="create")
def design_image(inp, sess):
    p = _new_file(inp.get("filename") or "Design", "." + (inp.get("format") or "png"))
    media.design_image(inp["spec"], p)
    mt, b64 = _image_block(p)
    return Result(f"Created {p}.", images=[(mt, b64)], files=[_file_rec(p)])


@tool("generate_image", "Make a photo-style or artistic picture from a description, using the image service set up in the Assistant settings. Write a rich, specific prompt (subject, style, light, composition). If no service is set up it says so; then use design_image instead.",
      {"prompt": {"type": "string"}, "size": {"type": "string", "description": "e.g. 1024x1024, 1536x1024, 1024x1536"}, "negative": {"type": "string"}, "filename": {"type": "string"}}, ["prompt"], risk="create")
def generate_image(inp, sess):
    p = _new_file(inp.get("filename") or inp["prompt"][:40], ".png")
    media.generate_image(inp["prompt"], p, inp.get("size") or "", inp.get("negative") or "")
    mt, b64 = _image_block(p)
    return Result(f"Created {p}.", images=[(mt, b64)], files=[_file_rec(p)])


@tool("text_to_speech", "Turn text into a spoken audio file (wav or mp3).", {"text": {"type": "string"}, "voice": {"type": "string"}, "filename": {"type": "string"}}, ["text"], risk="create")
def text_to_speech(inp, sess):
    p = _new_file(inp.get("filename") or "Speech", ".wav")
    r = media.speak(inp["text"], p, inp.get("voice") or "")
    return Result(f"Created {r['path']}.", files=[_file_rec(Path(r["path"]))])


@tool("make_video", "Make an .mp4 video from scenes: each scene {image | color (hex or [c1,c2]) | design (a design_image spec), title, text, caption, seconds, narration (spoken aloud with the voice), motion: zoom_in|zoom_out|pan_left|pan_right|none}. opts: width, height, fps, music (audio path), music_volume, fade, voice, narrate (bool). "
      "Pictures can come from generate_image, design_image or the user's files. Check a frame afterwards if the look matters.",
      {"scenes": {"type": "array", "items": {"type": "object"}}, "filename": {"type": "string"}, "options": {"type": "object"}}, ["scenes"], risk="create")
def make_video(inp, sess):
    p = _new_file(inp.get("filename") or "Video", ".mp4")
    r = media.make_video(inp["scenes"], p, dict(inp.get("options") or {}), None, sess.cancel if sess else None, (lambda f, m: sess.progress(m, f)) if sess else None)
    return Result(f"Created {p}: {r['seconds']} s, {r['scenes']} scenes, {r['size'][0]}x{r['size'][1]}.", files=[_file_rec(p)])


# ================================================================== the web
def _blocked_host(host: str) -> bool:
    import ipaddress
    import socket
    try:
        for fam, _t, _p, _c, sa in socket.getaddrinfo(host, None):
            ip = ipaddress.ip_address(sa[0])
            if ip.is_loopback or ip.is_link_local or ip.is_private or ip.is_reserved or ip.is_multicast:
                return True
    except Exception:
        return False
    return False


@tool("web_search", "Search the web (titles, links, snippets).", {"query": {"type": "string"}, "limit": {"type": "integer"}}, ["query"])
def web_search(inp, sess):
    if not store.load().get("web", True):
        raise UserError("Web access is turned off in the Assistant settings.")
    import requests
    from bs4 import BeautifulSoup
    r = requests.post("https://html.duckduckgo.com/html/", data={"q": inp["query"]}, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Fathom/2"}, timeout=25)
    s = BeautifulSoup(r.text, "html.parser")
    out = []
    for a in s.select("a.result__a")[: int(inp.get("limit") or 8)]:
        href = a.get("href", "")
        if "uddg=" in href:
            href = urllib.parse.unquote(href.split("uddg=", 1)[1].split("&", 1)[0])
        sn = a.find_parent(class_="result")
        snippet = sn.select_one(".result__snippet").get_text(" ", strip=True) if sn and sn.select_one(".result__snippet") else ""
        out.append(f"{a.get_text(' ', strip=True)}\n{href}\n{snippet}")
    return Result("\n\n".join(out) or "No results.")


@tool("web_fetch", "Read a web page (or PDF) as text.", {"url": {"type": "string"}, "max_chars": {"type": "integer"}}, ["url"])
def web_fetch(inp, sess):
    if not store.load().get("web", True):
        raise UserError("Web access is turned off in the Assistant settings.")
    import requests
    from bs4 import BeautifulSoup
    u = inp["url"].strip()
    host = urllib.parse.urlparse(u).hostname or ""
    if not u.startswith(("http://", "https://")) or _blocked_host(host):
        raise UserError("Only public web addresses can be opened.")
    r = requests.get(u, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Fathom/2"}, timeout=40)
    ct = r.headers.get("content-type", "")
    if "pdf" in ct or u.lower().endswith(".pdf"):
        from pypdf import PdfReader
        text = "\n".join((pg.extract_text() or "") for pg in PdfReader(io.BytesIO(r.content)).pages[:60])
    elif "html" in ct or "<html" in r.text[:500].lower():
        s = BeautifulSoup(r.text, "html.parser")
        for x in s(["script", "style", "nav", "footer", "noscript"]):
            x.decompose()
        text = re.sub(r"\n{3,}", "\n\n", s.get_text("\n"))
    else:
        text = r.text
    return Result(_clip(text, int(inp.get("max_chars") or 25000)))


# ================================================================== code
def _snapshot(d: Path) -> dict:
    out = {}
    try:
        for f in d.rglob("*"):
            if f.is_file():
                out[str(f)] = f.stat().st_mtime_ns
            if len(out) > 5000:
                break
    except OSError:
        pass
    return out


def _run_proc(cmd, sess, timeout: float, cwd: Path, shell=False) -> tuple[int, str]:
    env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
    p = subprocess.Popen(cmd, shell=shell, cwd=str(cwd), stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace", env=env, creationflags=no_window_flags())
    buf: list[str] = []

    def pump():
        for line in p.stdout:  # type: ignore[union-attr]
            buf.append(line)
            if len(buf) > 4000:
                del buf[:2000]
    th = threading.Thread(target=pump, daemon=True)
    th.start()
    t0 = time.time()
    while p.poll() is None:
        if sess is not None and sess.cancel.is_set():
            p.kill()
            return -1, "".join(buf) + "\n[stopped by the user]"
        if time.time() - t0 > timeout:
            p.kill()
            return -2, "".join(buf) + f"\n[stopped: ran longer than {int(timeout)} seconds]"
        time.sleep(0.1)
    th.join(2)
    return p.returncode, "".join(buf)


@tool("run_python", "Run Python code on the user's computer (Pillow, numpy, pandas-free basics, pypdf, reportlab, python-docx, openpyxl, cv2, matplotlib, requests are available). Use it for anything the other tools can't do: custom conversions, data crunching, batch edits, image processing, anything. The working folder is the assistant's output folder; new files there are reported. Print what you want to see.",
      {"code": {"type": "string"}, "timeout": {"type": "integer", "description": "seconds, default 300"}}, ["code"], risk="run")
def run_python(inp, sess):
    why = _blocked(inp["code"])
    if why:
        return Result("Refused: " + why + " Fathom never runs this, even when the assistant is set to never ask.", ok=False)
    d = store.output_dir()
    tmp = Path(tempfile.mkdtemp(prefix="fxpy_"))
    try:
        (tmp / "snippet.py").write_text(inp["code"], encoding="utf-8")
        before = _snapshot(d)
        rc, out = _run_proc([sys.executable, str(tmp / "snippet.py")], sess, float(inp.get("timeout") or 300), d)
        after = _snapshot(d)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    new = [Path(p) for p, m in after.items() if before.get(p) != m and "backups" not in Path(p).parts][:20]
    return Result(f"exit code {rc}\n" + _clip(out, 20000), files=[_file_rec(p) for p in new], ok=rc == 0)


_DANGER = [
    (r"\bformat(\.com)?\s+[a-z]:", "it would erase a whole drive."),
    (r"\bdiskpart\b|\bclean\s+all\b|\bmkfs(\.\w+)?\b|\bdd\s+if=.*of=/dev/", "it could wipe a disk or partition."),
    (r"(rm|rmdir|rd|del|erase|remove-item)\b[^\n|;&]*?(\s|^)(/s|-r|-rf|-fr|-recurse)\b[^\n|;&]*?(\s[a-z]:[\\/]?(\s|$|\*)|\s/(\s|$|\*)|\s~/?(\s|$)|windows|system32|program files|users[\\/]?(\s|$))", "it would delete a whole drive or system folder."),
    (r"\bcipher\s+/w", "it would overwrite free space on a drive."),
    (r"reg(\.exe)?\s+delete\s+hk(lm|cr|u)\\?\s*(/|$)|bcdedit\s+/(delete|set)|vssadmin\s+delete|wbadmin\s+delete|bootrec|sfc\s+/scannow\s*&&\s*del", "it would damage Windows itself."),
    (r"shutil\.rmtree\(\s*r?['\"](/|[a-z]:[\\/]?)['\"]", "it would delete a whole drive."),
    (r"\b(set-)?mpprefence\b|disable-windowsoptionalfeature|netsh\s+advfirewall\s+set\s+\w+\s+state\s+off|add-mppreference\s+-exclusion", "it would turn off Windows protection."),
]


def _blocked(text: str) -> str:
    if os.environ.get("FATHOM_SERVE") == "1" and os.environ.get("FATHOM_ALLOW_SHELL") != "1":
        return "this Fathom runs in server mode, where other devices can reach it, so commands and code are switched off (start it with FATHOM_ALLOW_SHELL=1 to allow them)."
    t = str(text or "").lower()
    for rx, why in _DANGER:
        if re.search(rx, t):
            return why
    return ""


@tool("run_shell", "Run a command line on the user's computer (PowerShell/cmd on Windows). For installing a package use: python -m pip install NAME. Be careful and explain what it does.",
      {"command": {"type": "string"}, "cwd": {"type": "string"}, "timeout": {"type": "integer"}}, ["command"], risk="run")
def run_shell(inp, sess):
    why = _blocked(inp["command"])
    if why:
        return Result("Refused: " + why + " Fathom never runs this, even when the assistant is set to never ask. Tell the user and offer a safer way.", ok=False)
    cwd = _resolve(inp["cwd"]) if inp.get("cwd") else store.output_dir()
    rc, out = _run_proc(inp["command"], sess, float(inp.get("timeout") or 300), cwd, shell=True)
    return Result(f"exit code {rc}\n" + _clip(out, 20000), ok=rc == 0)


# ================================================================== Fathom's own tools
@tool("fathom_tools", "List what Fathom can do: every screen (merge, split, rotate, stamp, convert, resize, OCR, redact, lock, rename, duplicates, hex view, drives, torrents ...) and the web routes behind them. Search with a few words, e.g. 'pdf', 'rotate', 'resize'. Then use fathom_tool_info and fathom_run.", {"query": {"type": "string"}})
def fathom_tools(inp, sess):
    return Result(json.dumps(catalog.listing(inp.get("query") or ""), ensure_ascii=False, indent=1)[:MAX_TEXT])


@tool("fathom_tool_info", "Exact details of one Fathom route: how its request body is read and the source of the function behind it (so you know every option). Always call this before fathom_run on a tool you have not used.", {"path": {"type": "string", "description": "e.g. /api/tools/rotate"}}, ["path"])
def fathom_tool_info(inp, sess):
    try:
        return Result(json.dumps(catalog.info(inp["path"]), ensure_ascii=False, indent=1)[:MAX_TEXT])
    except KeyError:
        raise UserError("Unknown route. Use fathom_tools to list them.")


def call_route(path: str, method: str, body: dict) -> object:
    """runs a Fathom web route in-process (no network)"""
    from app import main
    method = method.upper()
    for r in catalog.routes():
        pat = "^" + re.sub(r"\{(\w+)\}", r"(?P<\1>[^/]+)", r["path"]) + "$"
        m = re.match(pat, path)
        if m and r["method"] == method:
            fn = getattr(main, r["name"])
            hints = typing.get_type_hints(fn)
            kw = {}
            import inspect
            for name, prm in inspect.signature(fn).parameters.items():
                ann = hints.get(name)
                if name in m.groupdict():
                    kw[name] = m.group(name)
                elif isinstance(ann, type) and issubclass(ann, BaseModel):
                    kw[name] = ann(**{k: v for k, v in (body or {}).items()})
                elif name in (body or {}):
                    kw[name] = body[name]
                elif prm.default is inspect._empty:
                    if name == "request":
                        continue
                    raise UserError(f"This route needs '{name}' in the body.")
            return fn(**kw)
    raise UserError(f"No such route: {method} {path}")


@tool("fathom_run", "Run any Fathom tool. method POST/GET, path like /api/tools/rotate, body = the JSON the route reads (see fathom_tool_info). Routes that start a background job return {job}; this waits for it to finish and reports every result, including the files made. "
      "Results saved 'beside' the originals by default: unless the user wants otherwise pass body.output = {mode:'folder', dir:'<folder>'}. Proof pictures (before/after) are shown to the user automatically.",
      {"path": {"type": "string"}, "method": {"type": "string", "enum": ["POST", "GET"]}, "body": {"type": "object"}, "wait": {"type": "boolean"}}, ["path"], risk="write")
def fathom_run(inp, sess):
    path, method = inp["path"], (inp.get("method") or "POST").upper()
    body = dict(inp.get("body") or {})
    if "paths" in body:
        body["paths"] = [str(_resolve(x)) for x in body["paths"]]
    if body.get("output", {}).get("dir") if isinstance(body.get("output"), dict) else False:
        body["output"]["dir"] = str(_resolve(body["output"]["dir"]))
    res = call_route(path, method, body)
    if isinstance(res, dict) and isinstance(res.get("job"), str) and inp.get("wait", True):
        from app import ops, service
        jid = res["job"]
        is_search = path.startswith("/api/search")
        txt = _wait_job(sess, jid, search=is_search)
        proof, files = [], []
        job = ops.OPJOBS.get(jid)
        if job:
            for r in job.results:
                proof += (r.get("extra") or {}).get("proof") or []
                files += [{"path": o["path"], "name": o["name"], "kind": _kind(Path(o["path"])), "bytes": o["bytes"]} for o in r.get("outputs", [])[:6]]
        return Result(txt, files=files[:30], proof=proof[:12])
    return Result(_clip(json.dumps(_strip_big(res), ensure_ascii=False, default=str), 30000))


def _known_tool_ids() -> dict:
    from app.ai import agent
    tools = (agent.APP_MAP.get("tools") or [])
    if tools:
        return {t["id"]: t for t in tools}
    return {s["id"]: s for s in catalog.screens()} | {k: {"id": k, "label": k.title(), "group": "edit"} for k in ("stamp", "sign")}


@tool("find_in_fathom", "Where is something in the Fathom app? Searches the real screens (tabs and tools) and the places inside them, from a few plain words about what the user wants ('compress video', 'sign a pdf', 'dark mode', 'api key', 'drag arrows in a PDF'). "
      "Returns whether it exists, and for each match the tab, the tool, a sentence saying where to click, and the id to give to open_in_fathom. Use it for every 'where is / do you have / how do I find / can Fathom do' question about the app itself, and before open_in_fathom when unsure of the id.",
      {"query": {"type": "string", "description": "a few plain words describing what the user is looking for"}}, ["query"])
def find_in_fathom(inp, sess):
    from app.ai import agent
    return Result(json.dumps(catalog.find(str(inp.get("query") or ""), agent.APP_MAP.get("tools") or None, agent.APP_MAP.get("groups") or _GROUP_FALLBACK), ensure_ascii=False, indent=1)[:MAX_TEXT])


_GROUP_FALLBACK = {"home": "Home", "ai": "Assistant", "search": "Search", "convert": "Convert", "size": "Size", "edit": "Edit", "protect": "Protect", "tidy": "Tidy", "lab": "Lab", "sync": "Sync", "torrent": "Torrent", "drives": "Drives", "mine": "My tools"}


@tool("open_in_fathom", "TAKE THE USER THERE: switch the app to a tool (tab) by its id, e.g. 'visual', 'stamp', 'smartrename', 'aisettings' (ids come from find_in_fathom), and optionally add files to the left-hand list so they can continue by hand. "
      "When the user only asked WHERE something is, tell them and ask first; call this after they say yes, or straight away when they asked to be taken/opened/redirected.", {"tool": {"type": "string"}, "paths": {"type": "array", "items": {"type": "string"}}}, ["tool"])
def open_in_fathom(inp, sess):
    tid = str(inp.get("tool") or "").strip()
    known = _known_tool_ids()
    if tid.startswith("mine-") and tid[5:] in known:
        tid = tid[5:]
    if tid not in known and "mine-" + tid in known:
        tid = "mine-" + tid
    if tid not in known:
        close = catalog.find(tid.replace("-", " "), agent_tools(), _group_names(), limit=3)["matches"]
        if not close:
            names = {t.get("label", "").lower().replace(" ", ""): k for k, t in known.items()} | {k: k for k in known}
            close = [{"label": known[names[n]].get("label", n), "open_with": names[n]} for n in difflib.get_close_matches(tid.lower().replace(" ", ""), list(names), n=3, cutoff=0.6)]
        hint = "; ".join(f"{m['label']} (id {m['open_with']})" for m in {m["open_with"]: m for m in close if m.get("open_with")}.values())
        raise UserError(f"There is no screen called '{tid}'. " + (f"Did you mean: {hint}?" if hint else "Use find_in_fathom to look it up."))
    t = known[tid]
    grp = _group_names().get(t.get("group"), t.get("group") or "")
    return Result(f"Done: the app now shows {t.get('label') or tid}" + (f" (in the {grp} tab)" if grp else "") + ".", ui={"action": "open", "tool": tid, "label": t.get("label") or tid, "paths": [str(_resolve(x)) for x in inp.get("paths") or []]})


def agent_tools():
    from app.ai import agent
    return agent.APP_MAP.get("tools") or None


def _group_names() -> dict:
    from app.ai import agent
    return agent.APP_MAP.get("groups") or _GROUP_FALLBACK


# ================================================================== adding features to Fathom (My tools)
@tool("create_tool", "ADD A NEW FEATURE TO FATHOM. Creates a tool that appears in the app under 'My tools' with its own form, file list, where-to-save chooser, progress, results and hover tips. Use it when the user wants something Fathom does not do yet (or a custom workflow they will repeat). "
      "spec = {id (short slug), label, icon (one emoji), blurb (one plain sentence, shown as the hover tip), title, sub (explanation under the title), run (button text), mode: 'all' (run once with all files) | 'each' (once per file), accepts: any|pdf|image|video|audio|text|office, needs_files (bool), "
      "fields: [{id, type: text|number|textarea|select|check|range|choice|pages|password|datetime|note, label, tip (plain-words hover help for EVERY field), value (default), options: [{v, l, tip}], min, max, step}]}. "
      "code = the content of run.py. " + plugins.HARNESS_HELP + " After creating, call test_tool with a sample file to prove it works before telling the user it is ready.",
      {"spec": {"type": "object"}, "code": {"type": "string"}}, ["spec", "code"], risk="run")
def create_tool(inp, sess):
    meta = plugins.save(dict(inp["spec"]), inp["code"])
    return Result(f"Tool '{meta['label']}' saved (id {meta['id']}). It now appears in Fathom under My tools. Test it with test_tool.", ui={"action": "tools_changed", "open": meta["id"]})


@tool("test_tool", "Run one of the added tools once with sample files and inputs, and see exactly what it printed, made, or the error. Fix the code with create_tool (same id) until it works.", {"id": {"type": "string"}, "files": {"type": "array", "items": {"type": "string"}}, "inputs": {"type": "object"}}, ["id"], risk="run")
def test_tool(inp, sess):
    r = plugins.test_run(inp["id"], [str(_resolve(f)) for f in inp.get("files") or []], dict(inp.get("inputs") or {}), 900)
    files = [_file_rec(Path(p)) for p in r.get("outputs", []) if Path(p).exists()][:10]
    return Result(_clip(json.dumps(r, ensure_ascii=False, indent=1), 12000), files=files, ok=bool(r.get("ok")))


@tool("list_my_tools", "List the tools added to Fathom so far (id, label, what it does).", {})
def list_my_tools(inp, sess):
    return Result(json.dumps([{k: m[k] for k in ("id", "label", "blurb", "mode", "accepts")} for m in plugins.listing()], ensure_ascii=False, indent=1) or "[]")


@tool("read_my_tool", "Read the settings and code of an added tool (to improve it).", {"id": {"type": "string"}}, ["id"])
def read_my_tool(inp, sess):
    return Result(json.dumps(plugins.get(inp["id"]), ensure_ascii=False, indent=1)[:MAX_TEXT])


@tool("delete_my_tool", "Remove an added tool from Fathom.", {"id": {"type": "string"}}, ["id"], risk="run")
def delete_my_tool(inp, sess):
    plugins.delete(inp["id"])
    return Result(f"Removed {inp['id']}.", ui={"action": "tools_changed"})


# ================================================================== talking with the user
@tool("remember", "Save a fact or preference to remember in all future chats (e.g. the user's name, preferred folder, writing style, how they like documents). Keep it short.", {"note": {"type": "string"}, "forget": {"type": "boolean", "description": "remove this note instead"}}, ["note"])
def remember(inp, sess):
    mem = list(store.load().get("memory") or [])
    n = inp["note"].strip()
    if inp.get("forget"):
        mem = [m for m in mem if n.lower() not in m.lower()]
    elif n not in mem:
        mem.append(n)
    store.save({"memory": mem[-80:]})
    return Result("Saved." if not inp.get("forget") else "Forgotten.")


@tool("ask_user", "Ask the user a question and wait for the answer (use only when you truly cannot decide, or for a choice that matters). Offer 2–4 short choices when possible.", {"question": {"type": "string"}, "choices": {"type": "array", "items": {"type": "string"}}}, ["question"])
def ask_user(inp, sess):
    ans = sess.ask(inp["question"], inp.get("choices") or [])
    return Result("The user answered: " + ans)


@tool("show_files", "Show files to the user as clickable cards in the chat (every file you create is shown automatically; use this for other files).", {"paths": {"type": "array", "items": {"type": "string"}}}, ["paths"])
def show_files(inp, sess):
    ps = [_resolve(x) for x in inp["paths"]]
    return Result("Shown.", files=[_file_rec(p) for p in ps if p.exists()])


# ------------------------------------------------------------------ describing a call for the approval card
def describe(name: str, inp: dict) -> str:
    try:
        if name == "run_python":
            return "Run this Python code:\n" + str(inp.get("code", ""))[:2500]
        if name == "run_shell":
            return "Run this command:\n" + str(inp.get("command", ""))[:1500]
        if name == "write_file":
            return f"Write {inp.get('path')} ({len(str(inp.get('content', '')))} characters)"
        if name == "edit_file":
            return f"Edit {inp.get('path')}: replace\n{str(inp.get('old', ''))[:300]}\nwith\n{str(inp.get('new', ''))[:300]}"
        if name == "file_operation":
            return f"{inp.get('action')} {inp.get('path')}" + (f" → {inp.get('to')}" if inp.get("to") else "")
        if name == "fathom_run":
            b = inp.get("body") or {}
            return f"Run Fathom tool {inp.get('path')} on {len(b.get('paths') or [])} file(s)\n" + json.dumps({k: v for k, v in b.items() if k != "paths"}, ensure_ascii=False)[:800] + "\nFiles: " + ", ".join(map(str, (b.get("paths") or [])[:5]))
        if name == "create_tool":
            s = inp.get("spec") or {}
            return f"Add the tool '{s.get('label')}' to Fathom:\n{s.get('blurb', '')}\n\nCode:\n" + str(inp.get("code", ""))[:2500]
        if name == "test_tool":
            return f"Run the added tool '{inp.get('id')}' once as a test"
        if name == "delete_my_tool":
            return f"Remove the tool '{inp.get('id')}' from Fathom"
    except Exception:
        pass
    return name + " " + json.dumps(inp, ensure_ascii=False)[:600]
