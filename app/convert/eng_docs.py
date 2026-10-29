# Copyright (C) 2026 The Fathomist. Licensed under AGPL-3.0-or-later (see LICENSE).
"""Documents: PDF and e-book pages (PyMuPDF), text/markup (Pandoc), Microsoft Office / OpenDocument and old office formats
(LibreOffice), Kindle and other e-book formats (Calibre), and PDF -> Word with layout (pdf2docx)."""
from __future__ import annotations

import base64
import html as htmllib
import os
import re
import shutil
import subprocess
import sys
import textwrap
import threading
from functools import lru_cache
from pathlib import Path

from app.service import UserError
from app.convert.core import ConvertError, Ctx, Edge, Engine, Opts, find_program, parse_pages, safe_stem
from app.convert.formats import FORMATS
from app.runtime import no_window_flags


def get_fitz():
    try:
        import pymupdf as fitz  # type: ignore

        return fitz
    except ImportError:
        try:
            import fitz  # type: ignore

            return fitz
        except ImportError:
            return None


def read_text_file(path: Path) -> str:
    from app.extract import detect_encoding

    return path.read_text(encoding=detect_encoding(path), errors="replace")


# ---------------------------------------------------------------- PDF and page-based documents (PyMuPDF)
FITZ_DOCS = {"pdf", "xps", "epub", "mobi", "fb2", "cbz", "svg"}
PAGE_PICTURES = {"png", "jpg", "tif", "webp", "bmp", "avif", "jp2", "tga", "heic", "ppm"}


def open_doc(path: Path, o: Opts):
    fitz = get_fitz()
    try:
        doc = fitz.open(str(path))
    except Exception as exc:
        raise ConvertError(f"This document could not be opened ({exc}).")
    if getattr(doc, "needs_pass", False):
        if not o.password or not doc.authenticate(o.password):
            raise ConvertError("This PDF is password-protected. Enter the password in the options and try again.")
    return doc


class PageEngine(Engine):
    name = "pymupdf"
    label = "PDF and page-based documents"
    cost = 1.3

    def available(self) -> bool:
        return get_fitz() is not None

    def reads(self) -> set[str]:
        return FITZ_DOCS

    def writes(self) -> set[str]:
        from app.convert.eng_image import rw

        return (rw()[1] & PAGE_PICTURES) | {"pdf", "txt", "html"}

    def can(self, s: str, t: str) -> Edge | None:
        from app.convert.eng_image import rw

        if s not in FITZ_DOCS or s == t:
            return None
        imgs = rw()[1] & PAGE_PICTURES
        if t in imgs:
            return Edge(s, t, self, 1.3, "Every page becomes its own picture." if s != "svg" else "", multi=(s != "svg"))
        if t == "pdf":
            return Edge(s, t, self, 1.3)
        if t == "txt":
            return Edge(s, t, self, 1.4, "Only the text is kept (scanned pages are read with OCR when it is installed).")
        if t == "html":
            return Edge(s, t, self, 1.5, "The page layout is approximated.")
        return None

    def convert(self, src: Path, out_dir: Path, s: str, t: str, o: Opts, ctx: Ctx) -> list[Path]:
        fitz = get_fitz()
        out_dir.mkdir(parents=True, exist_ok=True)
        stem = safe_stem(src)
        doc = open_doc(src, o)
        try:
            if t == "pdf":
                dst = out_dir / f"{stem}.pdf"
                data = doc.convert_to_pdf()
                out = fitz.open("pdf", data)
                out.save(str(dst), garbage=3, deflate=True)
                out.close()
                return [dst]
            pages = parse_pages(o.pages, doc.page_count)
            if t == "txt":
                return [self._text(doc, pages, out_dir / f"{stem}.txt", o, ctx)]
            if t == "html":
                dst = out_dir / f"{stem}.html"
                body = []
                for n, pn in enumerate(pages, 1):
                    ctx.check()
                    body.append(doc[pn].get_text("html"))
                    ctx.progress(n / len(pages))
                dst.write_text("<!DOCTYPE html><html><head><meta charset='utf-8'><title>%s</title></head><body>%s</body></html>"
                               % (htmllib.escape(stem), "\n".join(body)), encoding="utf-8")
                return [dst]
            return self._images(doc, pages, out_dir, stem, t, o, ctx)
        finally:
            doc.close()

    def _images(self, doc, pages: list[int], out_dir: Path, stem: str, t: str, o: Opts, ctx: Ctx) -> list[Path]:
        from PIL import Image

        from app.convert.eng_image import save_pil

        outs: list[Path] = []
        alpha = t in ("png", "webp", "tga", "tif", "avif", "heic", "jp2", "dds", "icns", "sgi")
        dpi = max(20, min(1200, int(o.pdf_dpi or 150)))
        for n, pn in enumerate(pages, 1):
            ctx.check()
            page = doc[pn]
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            dst = out_dir / (f"{stem}.{t}" if len(pages) == 1 else f"{stem}_p{pn + 1:03d}.{t}")
            save_pil([img], [0], dst, t, o)
            outs.append(dst)
            ctx.progress(n / len(pages))
        return outs

    def _text(self, doc, pages: list[int], dst: Path, o: Opts, ctx: Ctx) -> Path:
        parts: list[str] = []
        ocr_used = 0
        for n, pn in enumerate(pages, 1):
            ctx.check()
            page = doc[pn]
            text = page.get_text("text").strip()
            if len(text) < 3 and page.get_images():
                text, used = self._ocr_page(page)
                ocr_used += used
            parts.append(text)
            ctx.progress(n / len(pages))
        result = "\n\n".join(parts).strip()
        if not result:
            raise ConvertError("This document has no text in it, and the OCR reader is not available to read its pictures.")
        dst.write_text(result + "\n", encoding="utf-8")
        return dst

    @staticmethod
    def _ocr_page(page) -> tuple[str, int]:
        try:
            from PIL import Image

            from app.convert.eng_image import lines_from_tokens
            from app.ocr import ocr_boxes
            from app.runtime import ML_LOCK

            pix = page.get_pixmap(dpi=200, alpha=False)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            with ML_LOCK:
                return lines_from_tokens(ocr_boxes(img)), 1
        except Exception:
            return "", 0


# ---------------------------------------------------------------- text and web pages -> PDF
def externalize_data_uris(html: str, folder: Path) -> str:
    """The PDF writer cannot read pictures embedded as data: addresses, so write them out as files."""
    count = [0]

    def repl(m: re.Match) -> str:
        count[0] += 1
        ext = {"jpeg": "jpg", "svg+xml": "svg"}.get(m.group(1).lower(), m.group(1).lower())
        name = f"_embedded_{count[0]}.{ext}"
        try:
            (folder / name).write_bytes(base64.b64decode(m.group(2)))
        except Exception:
            return m.group(0)
        return f'src="{name}"'

    return re.sub(r'src="data:image/([\w+.-]+);base64,([^"]+)"', repl, html)


def html_to_pdf(html: str, dst: Path, bases: list[Path], ctx: Ctx, css: str = "") -> None:
    fitz = get_fitz()
    base_css = "body{font-family:sans-serif;font-size:11pt;} img{max-width:100%;} table{border-collapse:collapse;} td,th{border:1px solid #888;padding:3px 6px;}"
    arch = None
    for b in bases:
        if arch is None:
            arch = fitz.Archive(str(b))
        else:
            try:
                arch.add(str(b))
            except Exception:
                pass
    story = fitz.Story(html=html, user_css=base_css + css, archive=arch)
    writer = fitz.DocumentWriter(str(dst))
    media = fitz.paper_rect("a4")
    where = media + (48, 48, -48, -48)
    more = 1
    pages = 0
    while more:
        ctx.check()
        dev = writer.begin_page(media)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
        pages += 1
        if pages > 20000:
            break
    writer.close()


class TextPdfEngine(Engine):
    name = "textpdf"
    label = "Text and web pages to PDF"
    cost = 1.3

    def available(self) -> bool:
        fitz = get_fitz()
        return bool(fitz and hasattr(fitz, "Story"))

    def reads(self) -> set[str]:
        return {"txt", "html"}

    def writes(self) -> set[str]:
        return {"pdf"}

    def convert(self, src: Path, out_dir: Path, s: str, t: str, o: Opts, ctx: Ctx) -> list[Path]:
        out_dir.mkdir(parents=True, exist_ok=True)
        dst = out_dir / f"{safe_stem(src)}.pdf"
        text = read_text_file(src)
        if s == "txt":
            lines: list[str] = []
            for line in text.replace("\t", "    ").splitlines():
                lines.extend(textwrap.wrap(line, 96, replace_whitespace=False, drop_whitespace=False) or [""])
            html = "<pre style='font-family:monospace;font-size:9pt;line-height:1.25'>%s</pre>" % htmllib.escape("\n".join(lines))
            html_to_pdf(html, dst, [], ctx)
        else:
            work = ctx.scratch_dir()
            html_to_pdf(externalize_data_uris(text, work), dst, [src.parent, work], ctx)
        return [dst]


# ---------------------------------------------------------------- Pandoc
@lru_cache(maxsize=1)
def pandoc_exe() -> str | None:
    try:
        import pypandoc  # type: ignore

        p = pypandoc.get_pandoc_path()
        if p and os.path.exists(p):
            return p
    except Exception:
        pass
    return find_program("pandoc", "pandoc.exe", extra=[os.path.join(os.environ.get("LOCALAPPDATA", ""), "Pandoc", "pandoc.exe"),
                                                       r"C:\Program Files\Pandoc\pandoc.exe"])


PANDOC_READ = {"docx": "docx", "odt": "odt", "rtf": "rtf", "html": "html", "md": "markdown", "rst": "rst", "tex": "latex", "epub": "epub",
               "txt": "html", "org": "org", "textile": "textile", "wiki": "mediawiki", "docbook": "docbook", "ipynb": "ipynb", "fb2": "fb2",
               "csv": "csv", "tsv": "tsv"}
PANDOC_WRITE = {"docx": "docx", "odt": "odt", "rtf": "rtf", "html": "html5", "md": "gfm", "rst": "rst", "tex": "latex", "epub": "epub3",
                "txt": "plain", "org": "org", "textile": "textile", "wiki": "mediawiki", "docbook": "docbook5", "asciidoc": "asciidoc",
                "ipynb": "ipynb", "fb2": "fb2", "pptx": "pptx", "odp": "odp"}


@lru_cache(maxsize=1)
def _pandoc_lists() -> tuple[set[str], set[str]]:
    exe = pandoc_exe()
    if not exe:
        return set(), set()

    def lst(flag: str) -> set[str]:
        try:
            out = subprocess.run([exe, flag], capture_output=True, text=True, timeout=30, creationflags=no_window_flags()).stdout
        except Exception:
            return set()
        return set(out.split())

    return lst("--list-input-formats"), lst("--list-output-formats")


class PandocEngine(Engine):
    name = "pandoc"
    label = "Documents and markup (Pandoc)"
    cost = 1.0
    needs = "Pandoc"

    def available(self) -> bool:
        return bool(pandoc_exe())

    def reads(self) -> set[str]:
        ins, _ = _pandoc_lists()
        return {k for k, v in PANDOC_READ.items() if v in ins or v == "html" or v == "markdown"}

    def writes(self) -> set[str]:
        _, outs = _pandoc_lists()
        return {k for k, v in PANDOC_WRITE.items() if v in outs}

    def can(self, s: str, t: str) -> Edge | None:
        e = super().can(s, t)
        if not e:
            return None
        if s in ("csv", "tsv") and t in ("docx", "odt", "html", "md", "rst", "tex", "epub", "txt", "pptx", "odp"):
            e.cost = 1.6   # prefer the spreadsheet engines for tables
        if t in ("txt", "md", "rst", "org", "textile", "wiki", "asciidoc") and s in ("docx", "odt", "epub", "html", "fb2", "rtf"):
            e.note = "Pictures are not included in this text format."
        if t in ("pptx", "odp"):
            e.note = "Slides are created from the headings of the document."
        if s == "txt":
            e.cost = 1.0
        word_like = {"docx", "doc", "odt", "rtf", "dotx", "dot", "ott", "docm"}
        if s in word_like and t in word_like:
            e.cost = 1.8   # Word <-> OpenDocument <-> RTF is LibreOffice's job: it keeps fonts, page setup, shapes, headers and tables where Pandoc rebuilds plain text
        return e

    def convert(self, src: Path, out_dir: Path, s: str, t: str, o: Opts, ctx: Ctx) -> list[Path]:
        out_dir.mkdir(parents=True, exist_ok=True)
        stem = safe_stem(src)
        dst = out_dir / f"{stem}.{t}"
        exe = pandoc_exe()
        work_src = src
        if s == "txt":  # plain text: keep every line break, do not guess at markdown
            text = read_text_file(src)
            paras = re.split(r"\n\s*\n", text.replace("\r\n", "\n"))
            body = "".join("<p>%s</p>\n" % htmllib.escape(p).replace("\n", "<br>\n") for p in paras if p.strip())
            work_src = ctx.scratch(".html")
            work_src.write_text("<html><head><meta charset='utf-8'><title>%s</title></head><body>%s</body></html>" % (htmllib.escape(stem), body), encoding="utf-8")
        cmd = [exe, str(work_src), "-f", PANDOC_READ[s], "-t", PANDOC_WRITE[t], "-o", str(dst), "--resource-path", str(src.parent)]
        if s in ("docx", "odt", "rtf", "epub", "fb2"):
            # a document's own Title (and Subtitle) is metadata to Pandoc and silently vanishes from the body: put it back as the first heading
            lua = ctx.scratch(".lua")
            lua.write_text("function Pandoc(doc)\n  local out = {}\n  local m = doc.meta\n"
                           "  if m.title then table.insert(out, pandoc.Header(1, {pandoc.Str(pandoc.utils.stringify(m.title))})) end\n"
                           "  if m.subtitle then table.insert(out, pandoc.Para({pandoc.Emph({pandoc.Str(pandoc.utils.stringify(m.subtitle))})})) end\n"
                           "  for _, b in ipairs(doc.blocks) do table.insert(out, b) end\n"
                           "  m.title = nil; m.subtitle = nil\n  return pandoc.Pandoc(out, m)\nend\n", encoding="utf-8")
            cmd += ["--lua-filter", str(lua)]
        if t in ("html", "tex", "docbook", "asciidoc", "org"):
            cmd += ["--standalone", "--metadata", f"pagetitle={stem}"] + ([] if s in ("docx", "odt", "rtf", "epub", "fb2") else ["--metadata", f"title={stem}"])
        elif t == "epub":
            cmd += ["--metadata", f"title={stem}"]
        if t == "html":
            cmd += ["--embed-resources", "--wrap=none"]
        elif t in ("md", "txt", "rst", "org", "textile", "wiki"):
            cmd += ["--wrap=none"]
        rc, out = ctx.run(cmd, timeout=900)
        if rc != 0 or not dst.exists():
            raise ConvertError("Pandoc could not convert this file: " + (out.splitlines()[-1] if out else f"exit code {rc}"))
        return [dst]


# ---------------------------------------------------------------- LibreOffice (Office formats and many old ones)
LO_LOCK = threading.Lock()
LO_WRITER = {"docx", "doc", "dotx", "dot", "odt", "ott", "rtf", "wpd", "wps", "sxw", "abw", "pages", "sdw", "stw", "hwp", "lwp", "t602", "mw", "zabw", "uot"}
LO_CALC = {"xlsx", "xls", "xlsm", "xltx", "ods", "sxc", "numbers", "dbf", "dif", "slk", "sdc", "stc", "lotus123", "wk1", "wks", "qpw", "wb2", "xlsb", "uos"}
LO_IMPRESS = {"pptx", "ppt", "ppsx", "pps", "potx", "odp", "sxi", "key", "sdd", "sti", "uop", "pot"}
LO_DRAW = {"vsdx", "vsd", "pub", "cdr", "odg", "vdx", "wpg", "cmx", "pmd", "qxp", "zmf", "dxf", "cgm", "svm", "fh"}
LO_WRITE_FROM = {
    "writer": {"pdf": "pdf", "docx": "docx", "doc": "doc", "odt": "odt", "rtf": "rtf", "txt": "txt:Text (encoded):UTF8", "html": "html",
               "epub": "epub", "png": "png", "jpg": "jpg", "ott": "ott", "dotx": "dotx"},
    "calc": {"pdf": "pdf", "xlsx": "xlsx", "xls": "xls", "ods": "ods", "csv": "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,-1",
             "html": "html", "png": "png", "jpg": "jpg", "xltx": "xltx", "dif": "dif", "dbf": "dbf", "slk": "slk",
             "tsv": "csv:Text - txt - csv (StarCalc):9,34,76,1,,0,false,true,false,false,false,-1"},
    "impress": {"pdf": "pdf", "pptx": "pptx", "ppt": "ppt", "odp": "odp", "png": "png", "jpg": "jpg", "html": "html", "gif": "gif", "bmp": "bmp", "tif": "tif", "svg": "svg",
                "potx": "potx", "ppsx": "ppsx", "pps": "pps"},
    "draw": {"pdf": "pdf", "png": "png", "jpg": "jpg", "svg": "svg", "odg": "odg", "gif": "gif", "bmp": "bmp", "tif": "tif", "emf": "emf", "wmf": "wmf", "eps": "eps"},
}
LO_FAMILY_IN = {"writer": LO_WRITER, "calc": LO_CALC, "impress": LO_IMPRESS, "draw": LO_DRAW}


@lru_cache(maxsize=1)
def soffice_exe() -> str | None:
    pf = [os.environ.get("ProgramFiles", r"C:\Program Files"), os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")]
    extra = [os.path.join(p, "LibreOffice", "program", "soffice.exe") for p in pf]
    extra += [os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "LibreOffice", "program", "soffice.exe"),
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    return find_program("soffice", "libreoffice", extra=extra)


def family_of(fmt: str) -> str | None:
    for fam, ids in LO_FAMILY_IN.items():
        if fmt in ids:
            return fam
    return None


class OfficeEngine(Engine):
    name = "office"
    label = "Microsoft Office, OpenDocument and old office formats (LibreOffice)"
    cost = 1.0
    heavy = True
    needs = "LibreOffice"

    def available(self) -> bool:
        return bool(soffice_exe())

    def reads(self) -> set[str]:
        return set().union(*LO_FAMILY_IN.values())

    def writes(self) -> set[str]:
        return set().union(*[set(v) for v in LO_WRITE_FROM.values()])

    def can(self, s: str, t: str) -> Edge | None:
        fam = family_of(s)
        if not fam or s == t or t not in LO_WRITE_FROM[fam]:
            return None
        note = ""
        multi = False
        if t == "png" or t == "jpg":
            note = "Only the first page / slide becomes a picture."
        if t == "csv":
            note = "Every sheet is saved as its own CSV file."
            multi = True
        if t == "html" and fam in ("calc", "impress"):
            note = "Layout is approximate."
        return Edge(s, t, self, 1.0 if t not in ("png", "jpg") else 3.0, note, multi=multi)

    def convert(self, src: Path, out_dir: Path, s: str, t: str, o: Opts, ctx: Ctx) -> list[Path]:
        exe = soffice_exe()
        fam = family_of(s)
        assert fam
        spec = LO_WRITE_FROM[fam][t]
        out_dir.mkdir(parents=True, exist_ok=True)
        from app.paths import data_root

        profile = data_root() / "lo_profile"
        profile.mkdir(parents=True, exist_ok=True)
        work = ctx.scratch_dir()
        # LibreOffice names the result after the input, so give it a simple copy-free name inside a scratch folder
        link = work / f"in_{safe_stem(src)[:60]}{src.suffix}"
        try:
            os.link(src, link)
        except OSError:
            shutil.copy2(src, link)
        ctx.status("Opening the document in LibreOffice…")
        with LO_LOCK:
            ctx.check()
            cmd = [exe, "--headless", "--norestore", "--nolockcheck", "--nodefault", "--nologo", "--nofirststartwizard",
                   f"-env:UserInstallation={profile.as_uri()}", "--convert-to", spec, "--outdir", str(work), str(link)]
            rc, out = ctx.run(cmd, timeout=1800)
        produced = sorted(p for p in work.iterdir() if p.is_file() and p.name != link.name)
        if not produced:
            raise ConvertError("LibreOffice could not convert this file" + (f": {out.splitlines()[-1]}" if out else "") +
                               ". It may be damaged or password-protected.")
        stem = safe_stem(src)
        outs: list[Path] = []
        if t == "html":
            main = next((p for p in produced if p.suffix.lower() in (".html", ".htm")), None)
            if main is None:
                raise ConvertError("LibreOffice did not produce a web page.")
            text = main.read_text(encoding="utf-8", errors="replace")
            for extra in produced:
                if extra is main:
                    continue
                mime = {".png": "png", ".jpg": "jpeg", ".jpeg": "jpeg", ".gif": "gif", ".svg": "svg+xml"}.get(extra.suffix.lower())
                if mime:
                    text = text.replace(extra.name, "data:image/%s;base64,%s" % (mime, base64.b64encode(extra.read_bytes()).decode()))
            main.write_text(text, encoding="utf-8")
            produced = [main]
        for p in produced:
            name = p.name
            tail = name[len("in_" + safe_stem(src)[:60]):] if name.startswith("in_" + safe_stem(src)[:60]) else "." + t
            if t == "csv" and len(produced) > 1:
                tail = tail.lstrip("-")  # "-SheetName.csv"
                dst = out_dir / f"{stem}_{Path(tail).stem}.csv" if Path(tail).stem not in ("", "csv") else out_dir / f"{stem}.csv"
            else:
                dst = out_dir / f"{stem}.{t}"
            shutil.move(str(p), str(dst))
            outs.append(dst)
        return outs


# ---------------------------------------------------------------- Calibre (Kindle and many e-book formats)
@lru_cache(maxsize=1)
def calibre_exe() -> str | None:
    pf = [os.environ.get("ProgramFiles", r"C:\Program Files"), os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")]
    extra = [os.path.join(p, "Calibre2", "ebook-convert.exe") for p in pf]
    extra += ["/Applications/calibre.app/Contents/MacOS/ebook-convert"]
    return find_program("ebook-convert", "ebook-convert.exe", extra=extra)


class CalibreEngine(Engine):
    name = "calibre"
    label = "E-books (Calibre)"
    cost = 1.2
    heavy = True
    needs = "Calibre"
    READ = {"epub", "mobi", "azw3", "azw", "fb2", "lit", "lrf", "pdb", "djvu", "cbz", "cbr", "htmlz", "txtz", "docx", "rtf", "txt", "html", "md", "odt", "pdf",
            "chm", "fbz", "pml", "rb", "snb", "tcr"}
    WRITE = {"epub", "mobi", "azw3", "fb2", "pdf", "docx", "txt", "htmlz", "txtz", "rtf", "lit", "lrf", "pdb", "pml", "rb", "snb", "tcr"}

    def available(self) -> bool:
        return bool(calibre_exe())

    def reads(self) -> set[str]:
        return self.READ

    def writes(self) -> set[str]:
        return self.WRITE

    def can(self, s: str, t: str) -> Edge | None:
        e = super().can(s, t)
        if e and s == "pdf":
            e.cost = 2.5
            e.note = "Converting a PDF to an e-book can lose layout."
        elif e and s in ("docx", "rtf", "txt", "html", "md", "odt") and t in ("docx", "rtf", "txt", "odt"):
            return None   # Pandoc and LibreOffice do these better
        return e

    def convert(self, src: Path, out_dir: Path, s: str, t: str, o: Opts, ctx: Ctx) -> list[Path]:
        out_dir.mkdir(parents=True, exist_ok=True)
        dst = out_dir / f"{safe_stem(src)}.{t}"
        # Calibre decides by file extension, so use a copy with the format's own extension
        work = ctx.scratch_dir()
        copy = work / f"book.{s}"
        try:
            os.link(src, copy)
        except OSError:
            shutil.copy2(src, copy)
        ctx.status("Converting the e-book with Calibre…")
        rc, out = ctx.run([calibre_exe(), copy, dst], timeout=3600)
        if rc != 0 or not dst.exists():
            raise ConvertError("Calibre could not convert this book: " + (out.splitlines()[-1] if out else f"exit code {rc}"))
        return [dst]


# ---------------------------------------------------------------- PDF -> Word keeping the layout
class PdfToWordEngine(Engine):
    """Diagrams, posters, forms and anything with boxes or pictures are rebuilt object by object so every shape, arrow, picture and line of text stays
    exactly where it was (convert/pdf_layout.py). Pages of running text are rebuilt as flowing Word text by pdf2docx when it is installed."""
    name = "pdf2docx"
    label = "PDF to Word with layout"
    cost = 1.2
    heavy = True
    needs = "python-docx"

    def available(self) -> bool:
        try:
            import importlib.util

            return importlib.util.find_spec("docx") is not None
        except Exception:
            return False

    def _flow_ok(self) -> bool:
        try:
            import importlib.util

            return importlib.util.find_spec("pdf2docx") is not None and get_fitz() is not None
        except Exception:
            return False

    def reads(self) -> set[str]:
        return {"pdf"}

    def writes(self) -> set[str]:
        return {"docx", "odt"}

    def can(self, s: str, t: str) -> Edge | None:
        e = super().can(s, t)
        if e:
            e.note = "Boxes, arrows, pictures and text keep their exact places (text stays editable). Plain text pages become flowing Word text."
            if t == "odt":
                e.cost = 1.1          # older Word / RTF routes start from this one: LibreOffice keeps its shapes where it drops those of a .docx
        return e

    def convert(self, src: Path, out_dir: Path, s: str, t: str, o: Opts, ctx: Ctx) -> list[Path]:
        from . import pdf_layout

        out_dir.mkdir(parents=True, exist_ok=True)
        dst = out_dir / f"{safe_stem(src)}.{t}"
        if o.password:
            raise ConvertError("Remove the password from this PDF first (Protect > PDF password).")
        if t == "odt":
            if self._flow_ok() and not pdf_layout.has_graphics(str(src)):
                raise ConvertError("flowing text: use the Word route")        # the next route (Word, then LibreOffice) reflows plain text properly
            pdf_layout.to_odt(str(src), dst, ctx.progress)
            return [dst]
        if pdf_layout.has_graphics(str(src)) or not self._flow_ok():
            try:
                pdf_layout.to_docx(str(src), dst, ctx.progress)
                return [dst]
            except UserError:
                raise
            except Exception as exc:
                if not self._flow_ok():
                    raise ConvertError(f"Could not rebuild this PDF as Word: {exc}")
        from pdf2docx import Converter  # type: ignore

        cv = Converter(str(src))
        try:
            cv.convert(str(dst), start=0, end=None)
        finally:
            cv.close()
        return [dst]


# ---------------------------------------------------------------- slides -> a readable document
class SlidesToDocEngine(Engine):
    """A presentation as a document: each slide's title, text (with its bullet levels), tables, pictures and speaker notes, slide by slide."""
    name = "slides2doc"
    label = "Slides to a document"
    cost = 0.8

    def available(self) -> bool:
        try:
            import importlib.util

            return importlib.util.find_spec("pptx") is not None and importlib.util.find_spec("docx") is not None
        except Exception:
            return False

    def reads(self) -> set[str]:
        return {"pptx", "ppsx", "ppt", "pps", "odp"}

    def writes(self) -> set[str]:
        return {"docx"}

    def can(self, s: str, t: str) -> Edge | None:
        e = super().can(s, t)
        if e:
            e.note = "Every slide's title, text, tables, pictures and speaker notes are listed slide by slide."
            if s in ("ppt", "pps", "odp") and not soffice_exe():
                return None
        return e

    def convert(self, src: Path, out_dir: Path, s: str, t: str, o: Opts, ctx: Ctx) -> list[Path]:
        import io

        import docx  # type: ignore
        from docx.shared import Inches  # type: ignore
        from pptx import Presentation  # type: ignore
        from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER  # type: ignore

        out_dir.mkdir(parents=True, exist_ok=True)
        work = src
        if s not in ("pptx", "ppsx"):
            work = OfficeEngine().convert(src, ctx.scratch_dir(), s, "pptx", o, ctx)[0] if s != "odp" else OfficeEngine().convert(src, ctx.scratch_dir(), s, "pptx", o, ctx)[0]
        prs = Presentation(str(work))
        d = docx.Document()
        d.core_properties.title = safe_stem(src)

        def flatten(shapes):
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from flatten(sh.shapes)
                else:
                    yield sh

        for n, slide in enumerate(prs.slides, 1):
            ctx.check()
            shapes = sorted(flatten(slide.shapes), key=lambda x: ((x.top or 0) // 91440, x.left or 0))
            title = None
            try:
                title = slide.shapes.title
            except Exception:
                title = None
            ttext = title.text_frame.text.strip() if title is not None and title.has_text_frame else ""
            d.add_heading(f"Slide {n}" + (f": {ttext}" if ttext else ""), level=1)
            for sh in shapes:
                if title is not None and sh.shape_id == title.shape_id:
                    continue
                if sh.has_text_frame and sh.text_frame.text.strip():
                    body = sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
                    for para in sh.text_frame.paragraphs:
                        txt = "".join(r.text for r in para.runs).strip()
                        if not txt:
                            continue
                        lvl = min(2, int(para.level or 0))
                        if body or lvl:
                            p = d.add_paragraph(style="List Bullet" if lvl == 0 else f"List Bullet {lvl + 1}")
                        else:
                            p = d.add_paragraph()
                        for r in para.runs:
                            rr = p.add_run(r.text)
                            rr.bold = r.font.bold
                            rr.italic = r.font.italic
                elif getattr(sh, "has_table", False) and sh.has_table:
                    tb = sh.table
                    t_ = d.add_table(rows=len(tb.rows), cols=len(tb.columns))
                    t_.style = "Table Grid"
                    for i, row in enumerate(tb.rows):
                        for j, cell in enumerate(row.cells):
                            t_.cell(i, j).text = cell.text
                elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        w = min(6.0, max(1.0, (sh.width or 0) / 914400))
                        d.add_picture(io.BytesIO(sh.image.blob), width=Inches(w))
                    except Exception:
                        pass
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                    p = d.add_paragraph()
                    r = p.add_run("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
                    r.italic = True
            except Exception:
                pass
            ctx.progress(n / len(prs.slides))
        dst = out_dir / f"{safe_stem(src)}.docx"
        d.save(str(dst))
        return [dst]


# ---------------------------------------------------------------- PDF tables -> Excel / CSV, PDF pages -> PowerPoint
def _clean_cell(v) -> str:
    return "" if v is None else str(v).replace("\r", " ").strip()


_NUM = re.compile(r"^-?(?:0|[1-9]\d{0,2}(?:,\d{3})+|[1-9]\d*)(?:\.\d+)?$|^-?0\.\d+$")


def _as_number(c: str):
    """'1,234.50' -> 1234.5 so Excel can add it up; ids like '00123' or '+92...' stay text."""
    if c and _NUM.match(c):
        try:
            return float(c.replace(",", "")) if ("." in c) else int(c.replace(",", ""))
        except ValueError:
            return c
    return c


class PdfTablesEngine(Engine):
    name = "pdftables"
    label = "PDF tables to Excel / CSV"
    cost = 1.6

    def available(self) -> bool:
        return get_fitz() is not None

    def reads(self) -> set[str]:
        return {"pdf"}

    def writes(self) -> set[str]:
        return {"xlsx", "csv"}

    def can(self, s: str, t: str) -> Edge | None:
        e = super().can(s, t)
        if e:
            e.note = "Every table Fathom can find in the PDF is copied out (xlsx: one sheet per table; csv: one file per table). Scanned PDFs have no real tables."
            e.multi = t == "csv"
        return e

    def convert(self, src: Path, out_dir: Path, s: str, t: str, o: Opts, ctx: Ctx) -> list[Path]:
        out_dir.mkdir(parents=True, exist_ok=True)
        stem = safe_stem(src)
        doc = open_doc(src, o)
        tables: list[tuple[int, int, list[list[str]]]] = []
        try:
            pages = parse_pages(o.pages, doc.page_count)
            for n, pn in enumerate(pages, 1):
                ctx.check()
                try:
                    found = doc[pn].find_tables()
                    items = list(getattr(found, "tables", found))
                except Exception:
                    items = []
                for k, tb in enumerate(items, 1):
                    try:
                        rows = [[_clean_cell(c) for c in row] for row in tb.extract()]
                    except Exception:
                        continue
                    rows = [r for r in rows if any(r)]
                    if rows:
                        tables.append((pn + 1, k, rows))
                ctx.progress(n / len(pages) * 0.8)
        finally:
            doc.close()
        if not tables:
            raise ConvertError("No tables were found in this PDF. (Only PDFs with real text tables work; a scanned page must be turned into text first.)")
        if t == "csv":
            import csv

            outs: list[Path] = []
            for pg, k, rows in tables:
                dst = out_dir / (f"{stem}.csv" if len(tables) == 1 else f"{stem}_p{pg:03d}_t{k}.csv")
                with dst.open("w", newline="", encoding="utf-8-sig") as fh:
                    csv.writer(fh).writerows(rows)
                outs.append(dst)
            ctx.progress(1.0)
            return outs
        import openpyxl

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for pg, k, rows in tables:
            ws = wb.create_sheet(f"p{pg}_t{k}"[:31])
            for r in rows:
                cells = [_as_number(c) for c in r]
                ws.append(cells)
        dst = out_dir / f"{stem}.xlsx"
        wb.save(str(dst))
        ctx.progress(1.0)
        return [dst]


class PdfSlidesEngine(Engine):
    name = "pdfslides"
    label = "PDF pages to PowerPoint"
    cost = 1.7

    def available(self) -> bool:
        try:
            import importlib.util

            return importlib.util.find_spec("pptx") is not None
        except Exception:
            return False

    def reads(self) -> set[str]:
        return {"pdf"}

    def writes(self) -> set[str]:
        return {"pptx"}

    def can(self, s: str, t: str) -> Edge | None:
        e = super().can(s, t)
        if e:
            e.note = "Each page becomes one slide; boxes, arrows, pictures and text are rebuilt as editable objects in the same places (a scanned page stays a picture)."
        return e

    def convert(self, src: Path, out_dir: Path, s: str, t: str, o: Opts, ctx: Ctx) -> list[Path]:
        import io

        from pptx import Presentation  # type: ignore
        from pptx.util import Emu  # type: ignore

        out_dir.mkdir(parents=True, exist_ok=True)
        dst = out_dir / f"{safe_stem(src)}.pptx"
        if not o.pages and not o.password:
            try:
                from . import pdf_layout

                pdf_layout.to_pptx(str(src), dst, ctx.progress)
                return [dst]
            except UserError:
                raise
            except Exception:
                if get_fitz() is None:
                    raise ConvertError("Could not rebuild this PDF as slides.")
        doc = open_doc(src, o)
        try:
            pages = parse_pages(o.pages, doc.page_count)
            prs = Presentation()
            first = doc[pages[0]].rect
            prs.slide_width = Emu(int(first.width * 12700))
            prs.slide_height = Emu(int(first.height * 12700))
            blank = prs.slide_layouts[6]
            dpi = max(72, min(300, int(o.pdf_dpi or 150)))
            for n, pn in enumerate(pages, 1):
                ctx.check()
                page = doc[pn]
                pix = page.get_pixmap(dpi=dpi, alpha=False)
                slide = prs.slides.add_slide(blank)
                w = prs.slide_width
                h = int(w * pix.height / max(1, pix.width))
                if h > prs.slide_height:
                    h = prs.slide_height
                    w = int(h * pix.width / max(1, pix.height))
                left = int((prs.slide_width - w) / 2)
                top = int((prs.slide_height - h) / 2)
                slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), left, top, width=w, height=h)
                try:
                    text = page.get_text("text").strip()
                    if text:
                        slide.notes_slide.notes_text_frame.text = text
                except Exception:
                    pass
                ctx.progress(n / len(pages))
            prs.save(str(dst))
        finally:
            doc.close()
        return [dst]
