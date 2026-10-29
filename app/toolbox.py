# Copyright (C) 2026 The Fathomist. Licensed under AGPL-3.0-or-later (see LICENSE).
"""Toolbox: barcode maker, PDF watermark remover, searchable PDF (OCR), background remover, subtitle fixer and burner,
and the quick text helpers. All of it runs on this computer."""
from __future__ import annotations

import base64
import io
import re
import shutil
import subprocess
from collections import Counter
from pathlib import Path

from app import ops
from app.convert.core import ConvertError, Ctx, ffmpeg_exe
from app.ops import Output, OutputPolicy, Task
from app.opsvc import expand_sources
from app.service import UserError


def _int(v, d, lo, hi):
    try:
        n = int(float(v))
    except (TypeError, ValueError):
        n = d
    return max(lo, min(hi, n))


def _rgb(s, d=(0, 0, 0)):
    s = str(s or "").strip().lstrip("#")
    if len(s) == 3:
        s = "".join(c * 2 for c in s)
    try:
        return int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16)
    except ValueError:
        return d


# ================================================================ barcodes
C128 = ("212222 222122 222221 121223 121322 131222 122213 122312 132212 221213 221312 231212 112232 122132 122231 113222 123122 123221 223211 221132 "
        "221231 213212 223112 312131 311222 321122 321221 312212 322112 322211 212123 212321 232121 111323 131123 131321 112313 132113 132311 211313 "
        "231113 231311 112133 112331 132131 113123 113321 133121 313121 211331 231131 213113 213311 213131 311123 311321 331121 312113 312311 332111 "
        "314111 221411 431111 111224 111422 121124 121421 141122 141221 112214 112412 122114 122411 142112 142211 241211 221114 413111 241112 134111 "
        "111242 121142 121241 114212 124112 124211 411212 421112 421211 212141 214121 412121 111143 111341 131141 114113 114311 411113 411311 113141 "
        "114131 311141 411131 211412 211214 211232 2331112").split()
C39 = {"0": "nnnwwnwnn", "1": "wnnwnnnnw", "2": "nnwwnnnnw", "3": "wnwwnnnnn", "4": "nnnwwnnnw", "5": "wnnwwnnnn", "6": "nnwwwnnnn", "7": "nnnwnnwnw", "8": "wnnwnnwnn", "9": "nnwwnnwnn",
       "A": "wnnnnwnnw", "B": "nnwnnwnnw", "C": "wnwnnwnnn", "D": "nnnnwwnnw", "E": "wnnnwwnnn", "F": "nnwnwwnnn", "G": "nnnnnwwnw", "H": "wnnnnwwnn", "I": "nnwnnwwnn", "J": "nnnnwwwnn",
       "K": "wnnnnnnww", "L": "nnwnnnnww", "M": "wnwnnnnwn", "N": "nnnnwnnww", "O": "wnnnwnnwn", "P": "nnwnwnnwn", "Q": "nnnnnnwww", "R": "wnnnnnwwn", "S": "nnwnnnwwn", "T": "nnnnwnwwn",
       "U": "wwnnnnnnw", "V": "nwwnnnnnw", "W": "wwwnnnnnn", "X": "nwnnwnnnw", "Y": "wwnnwnnnn", "Z": "nwwnwnnnn", "-": "nwnnnnwnw", ".": "wwnnnnwnn", " ": "nwwnnnwnn", "*": "nwnnwnwnn",
       "$": "nwnwnwnnn", "/": "nwnwnnnwn", "+": "nwnnnwnwn", "%": "nnnwnwnwn"}
EAN_L = "0001101 0011001 0010011 0111101 0100011 0110001 0101111 0111011 0110111 0001011".split()
EAN_G = [("".join("1" if c == "0" else "0" for c in x))[::-1] for x in EAN_L]
EAN_R = ["".join("1" if c == "0" else "0" for c in x) for x in EAN_L]
EAN_PAR = "LLLLLL LLGLGG LLGGLG LLGGGL LGLLGG LGGLLG LGGGLL LGLGLG LGLGGL LGGLGL".split()
ITF = ["nnwwn", "wnnnw", "nwnnw", "wwnnn", "nnwnw", "wnwnn", "nwwnn", "nnnww", "wnnwn", "nwnwn"]

KINDS = {"code128": "Code 128 (any text)", "code39": "Code 39 (capitals, digits, - . $ / + %)", "ean13": "EAN-13 (shop products, 12 or 13 digits)", "ean8": "EAN-8 (small products, 7 or 8 digits)",
         "upca": "UPC-A (US products, 11 or 12 digits)", "itf": "Interleaved 2 of 5 (digits, cartons)"}


def _check_digit(d: str) -> str:
    s = sum(int(c) * (3 if (len(d) - i) % 2 else 1) for i, c in enumerate(d))
    return str((10 - s % 10) % 10)


def _widths_from_bits(bits: str) -> list[int]:
    out, cur, n = [], bits[0], 0
    for c in bits:
        if c == cur:
            n += 1
        else:
            out.append(n); cur = c; n = 1
    out.append(n)
    return out if bits[0] == "1" else [0] + out      # always start with a bar


def encode(kind: str, text: str) -> tuple[list[int], str]:
    """Returns (alternating bar/space widths starting with a bar, the text to print under it)."""
    if not text:
        raise UserError("Type something to turn into a barcode.")
    if kind == "code128":
        try:
            text.encode("latin-1")
        except UnicodeEncodeError:
            raise UserError("Code 128 handles letters, digits and symbols found on a normal keyboard. Remove accents or other scripts, or use a QR code.")
        vals = []
        digits = text.isdigit() and len(text) >= 4 and len(text) % 2 == 0
        if digits:
            start, vals = 105, [int(text[i:i + 2]) for i in range(0, len(text), 2)]
        else:
            start = 104
            for ch in text:
                if not 32 <= ord(ch) <= 126:
                    raise UserError("Code 128 here supports printable characters only (no tabs, no new lines).")
                vals.append(ord(ch) - 32)
        chk = (start + sum(v * (i + 1) for i, v in enumerate(vals))) % 103
        seq = [start] + vals + [chk, 106]
        return [int(c) for v in seq for c in C128[v]], text
    if kind == "code39":
        t = text.upper()
        bad = sorted({c for c in t if c not in C39 or c == "*"})
        if bad:
            raise UserError("Code 39 cannot hold: " + " ".join(bad) + ". It allows capitals, digits, space and - . $ / + %")
        w: list[int] = []
        for c in "*" + t + "*":
            p = C39[c]
            for i, x in enumerate(p):
                w.append(3 if x == "w" else 1)
            w.append(1)
        return w[:-1], t
    if kind in ("ean13", "ean8", "upca"):
        if not text.isdigit():
            raise UserError("This barcode type takes digits only.")
        if kind == "upca":
            if len(text) not in (11, 12):
                raise UserError("UPC-A needs 11 digits (the 12th is added) or all 12.")
            body = text[:11]
            full = body + _check_digit(body)
            if len(text) == 12 and text != full:
                raise UserError(f"The last digit should be {full[-1]} for these numbers.")
            text, kind = "0" + full, "ean13"
            shown = full
        elif kind == "ean13":
            if len(text) not in (12, 13):
                raise UserError("EAN-13 needs 12 digits (the 13th is added) or all 13.")
            full = text[:12] + _check_digit(text[:12])
            if len(text) == 13 and text != full:
                raise UserError(f"The last digit should be {full[-1]} for these numbers.")
            text = shown = full
        else:
            if len(text) not in (7, 8):
                raise UserError("EAN-8 needs 7 digits (the 8th is added) or all 8.")
            full = text[:7] + _check_digit(text[:7])
            if len(text) == 8 and text != full:
                raise UserError(f"The last digit should be {full[-1]} for these numbers.")
            bits = "101" + "".join(EAN_L[int(c)] for c in full[:4]) + "01010" + "".join(EAN_R[int(c)] for c in full[4:]) + "101"
            return _widths_from_bits(bits), full
        par = EAN_PAR[int(text[0])]
        left = "".join((EAN_L if par[i] == "L" else EAN_G)[int(text[1 + i])] for i in range(6))
        bits = "101" + left + "01010" + "".join(EAN_R[int(c)] for c in text[7:]) + "101"
        return _widths_from_bits(bits), shown
    if kind == "itf":
        if not text.isdigit():
            raise UserError("Interleaved 2 of 5 takes digits only.")
        if len(text) % 2:
            text = "0" + text
        w = [1, 1, 1, 1]
        for i in range(0, len(text), 2):
            a, b = ITF[int(text[i])], ITF[int(text[i + 1])]
            for x, y in zip(a, b):
                w += [3 if x == "w" else 1, 3 if y == "w" else 1]
        w += [3, 1, 1]
        return w, text
    raise UserError("Unknown barcode type.")


def render_barcode(kind: str, text: str, o: dict):
    """Returns (PIL image, svg text)."""
    from PIL import Image, ImageDraw, ImageFont
    widths, shown = encode(kind, text)
    scale = _int(o.get("scale"), 3, 1, 20)
    height = _int(o.get("height"), 90, 20, 600)
    quiet = _int(o.get("quiet"), 10, 0, 40)
    fg, bg = _rgb(o.get("fg"), (0, 0, 0)), _rgb(o.get("bg"), (255, 255, 255))
    label = bool(o.get("label", True))
    pad = _int(o.get("margin"), 8, 0, 100)
    fs = _int(o.get("fontsize"), 14, 8, 60)
    total = sum(widths)
    W = (total + 2 * quiet) * scale + 2 * pad
    lab_h = fs + 8 if label else 0
    H = height + lab_h + 2 * pad
    img = Image.new("RGB", (W, H), bg)
    d = ImageDraw.Draw(img)
    x = pad + quiet * scale
    rects = []
    for i, w in enumerate(widths):
        if i % 2 == 0:
            d.rectangle([x, pad, x + w * scale - 1, pad + height - 1], fill=fg)
            rects.append((x, w * scale))
        x += w * scale
    font = None
    if label:
        try:
            font = ImageFont.load_default(size=fs)
        except TypeError:
            font = ImageFont.load_default()
        tw = d.textlength(shown, font=font)
        d.text(((W - tw) / 2, pad + height + 3), shown, fill=fg, font=font)
    hexf = "#%02x%02x%02x" % fg
    hexb = "#%02x%02x%02x" % bg
    svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">', f'<rect width="100%" height="100%" fill="{hexb}"/>']
    svg += [f'<rect x="{rx}" y="{pad}" width="{rw}" height="{height}" fill="{hexf}"/>' for rx, rw in rects]
    if label:
        esc = shown.replace("&", "&amp;").replace("<", "&lt;")
        svg.append(f'<text x="{W / 2}" y="{pad + height + fs + 1}" font-family="monospace" font-size="{fs}" text-anchor="middle" fill="{hexf}">{esc}</text>')
    svg.append("</svg>")
    return img, "\n".join(svg)


def barcode_preview(kind: str, text: str, o: dict) -> dict:
    line = next((l.strip() for l in text.splitlines() if l.strip()), "")
    img, _ = render_barcode(kind, line, o)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    _, shown = encode(kind, line)
    return {"image": "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode(), "shown": shown}


def barcode_start(kind: str, text: str, o: dict, output: dict, anchor: str = "") -> dict:
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if not lines:
        raise UserError("Type at least one value (one per line).")
    if len(lines) > 500:
        raise UserError("Up to 500 barcodes at a time.")
    for l in lines:
        encode(kind, l)                    # fail early with a clear message
    if not (output or {}).get("dir") or (output or {}).get("mode") != "folder":
        output = {"mode": "folder", "dir": str(Path.home() / "Documents" / "Fathom barcodes")}
    policy = OutputPolicy.from_dict(output)
    fmts = [f for f in (o.get("formats") or ["png"]) if f in ("png", "svg", "jpg", "pdf")] or ["png"]
    base_src = Path(anchor) if anchor and Path(anchor).exists() else None
    tasks = []
    for k, l in enumerate(lines):
        def fn(ctx: Ctx, task: Task, l=l, k=k) -> Output:
            img, svg = render_barcode(kind, l, o)
            d = ctx.scratch_dir()
            stem = re.sub(r"[^\w.-]+", "_", l)[:60] or f"barcode{k + 1}"
            made = []
            for f in fmts:
                p = d / f"{stem}.{f}"
                if f == "svg":
                    p.write_text(svg, encoding="utf-8")
                elif f == "jpg":
                    img.save(p, "JPEG", quality=95)
                elif f == "pdf":
                    img.save(p, "PDF", resolution=_int(o.get("dpi"), 150, 72, 600))
                else:
                    img.save(p, "PNG")
                made.append(p)
            return Output(policy.place(made, d / "x", Path(".")), "")
        tasks.append(Task(l[:40], fn, Path(l), rel=Path(".")))
    return {"job": ops.start("barcode", "Making barcodes", tasks).id, "count": len(tasks)}


# ================================================================ PDF: watermark remover
def _wm_one(src: Path, dst: Path, o: dict) -> str:
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    doc = fitz.open(str(src))
    if doc.needs_pass:
        raise ConvertError("This PDF is locked with a password. Remove the password first (Protect → PDF password).")
    mode = o.get("mode", "auto")
    n = len(doc)
    removed = Counter()
    pages = range(n)
    words = [w.strip() for w in re.split(r"[\n,;]+", o.get("words", "")) if w.strip()]
    if mode in ("words",) and not words:
        raise UserError("Type the words of the watermark to remove (one per line).")
    case = bool(o.get("case"))

    def redact(page, rects):
        for r in rects:
            page.add_redact_annot(r, fill=False)
        if rects:
            page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE) if hasattr(fitz, "PDF_REDACT_LINE_ART_NONE") else page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE)

    if o.get("annots", True) and mode in ("auto", "annots", "words"):
        for page in doc:
            for a in list(page.annots() or []):
                if a.type[1] in ("Watermark", "Stamp") or (o.get("freetext") and a.type[1] == "FreeText"):
                    page.delete_annot(a); removed["stamps and watermark notes"] += 1
    if o.get("layers", True) and mode in ("auto", "annots", "words"):
        try:
            ocgs = doc.get_ocgs() or {}
            hide = [x for x, v in ocgs.items() if re.search(r"water|draft|confidential|sample|stamp", v.get("name", ""), re.I)]
            if hide:
                doc.set_layer(-1, off=hide); removed["hidden layers"] += len(hide)
        except Exception:
            pass
    if mode == "words":
        for page in doc:
            rects = []
            for w in words:
                found = page.search_for(w) if case else page.search_for(w)
                rects += found
            redact(page, rects); removed["pieces of text"] += len(rects)
    if mode == "auto":
        thresh = max(2, int(n * _int(o.get("share"), 60, 20, 100) / 100 + 0.5)) if n > 1 else 1
        seen: dict = Counter()
        where: dict = {}
        for pi, page in enumerate(doc):
            for b in page.get_text("blocks"):
                if b[6] != 0:
                    continue
                t = re.sub(r"\s+", " ", b[4]).strip()
                if not t or t.isdigit():
                    continue
                key = (t.lower(), round(b[0] / 12), round(b[1] / 12))
                if (pi, key) in where:
                    continue
                where[(pi, key)] = fitz.Rect(b[:4])
                seen[key] += 1
        repeated = {k for k, c in seen.items() if c >= thresh and n > 1}
        if o.get("skiptop", True):
            ph = doc[0].rect.height
            repeated = {k for k in repeated if 0.1 * ph < k[2] * 12 < 0.88 * ph or len(k[0]) > 40 or True}
        if n == 1:
            repeated = set()
        for pi, page in enumerate(doc):
            rects = [where[(pi, k)] for k in repeated if (pi, k) in where]
            if rects:
                redact(page, rects); removed["repeated texts"] += len(rects)
        # pictures repeated on most pages (logos, big faint stamps)
        if o.get("images", True) and n > 1:
            cnt = Counter()
            for page in doc:
                for im in page.get_images(full=True):
                    cnt[im[0]] += 1
            big = {x for x, c in cnt.items() if c >= thresh}
            for page in doc:
                for x in big:
                    try:
                        page.delete_image(x); removed["repeated pictures"] += 1
                    except Exception:
                        pass
    if mode == "images":
        minshare = _int(o.get("share"), 60, 10, 100)
        cnt = Counter()
        for page in doc:
            for im in page.get_images(full=True):
                cnt[im[0]] += 1
        big = {x for x, c in cnt.items() if c >= max(1, int(n * minshare / 100 + 0.5))}
        for page in doc:
            for x in big:
                try:
                    page.delete_image(x); removed["pictures"] += 1
                except Exception:
                    pass
    if not removed:
        doc.close()
        raise ConvertError("Fathom did not find a watermark to remove. Try 'By words' and type the watermark text.")
    doc.save(str(dst), garbage=3, deflate=True)
    doc.close()
    return "Removed " + ", ".join(f"{v} {k}" for k, v in removed.items())


def wm_start(paths, recursive, o, output):
    policy = OutputPolicy.from_dict(output)
    tasks = []
    for src, rel in expand_sources(paths, recursive):
        if src.suffix.lower() != ".pdf":
            continue

        def fn(ctx, task, src=src, rel=rel):
            dst = ctx.scratch_dir() / f"{src.stem} (no watermark).pdf"
            note = _wm_one(src, dst, o)
            return Output(policy.place([dst], src, rel), note)
        tasks.append(Task(src.name, fn, src, rel=rel))
    if not tasks:
        raise UserError("Choose one or more PDF files.")
    return {"job": ops.start("watermark", "Removing watermarks", tasks).id, "count": len(tasks)}


# ================================================================ searchable PDF
def _ocr_lines(img):
    try:
        from app.ocr import ocr_boxes
        import app.service as service
        if service._has("rapidocr_onnxruntime"):
            return [(t["text"], t["bbox"]) for t in ocr_boxes(img) if t["text"].strip()]
    except Exception:
        pass
    try:
        import pytesseract
        d = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
        lines: dict = {}
        for i, t in enumerate(d["text"]):
            if t.strip() and float(d["conf"][i]) > 0:
                k = (d["block_num"][i], d["par_num"][i], d["line_num"][i])
                x, y, w, h = d["left"][i], d["top"][i], d["width"][i], d["height"][i]
                s = lines.setdefault(k, [[], x, y, x + w, y + h])
                s[0].append(t); s[1] = min(s[1], x); s[2] = min(s[2], y); s[3] = max(s[3], x + w); s[4] = max(s[4], y + h)
        return [(" ".join(v[0]), (v[1], v[2], v[3], v[4])) for v in lines.values()]
    except Exception:
        raise ConvertError("Fathom's text reader is not installed. Open Settings → Add-ons and install 'Read text from pictures'.")


def ocr_pdf_one(ctx, src: Path, dst: Path, o: dict) -> str:
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    from PIL import Image
    doc = fitz.open(str(src))
    if doc.needs_pass:
        raise ConvertError("This PDF is locked with a password. Remove the password first.")
    dpi = _int(o.get("dpi"), 200, 100, 400)
    skip = bool(o.get("skip", True))
    out = fitz.open()
    done = empty = 0
    for pi, page in enumerate(doc):
        ctx.check()
        has_text = len(page.get_text().strip()) > 20
        if skip and has_text:
            out.insert_pdf(doc, from_page=pi, to_page=pi)
            empty += 1
            continue
        pix = page.get_pixmap(dpi=dpi)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        lines = _ocr_lines(img)
        newp = out.new_page(width=page.rect.width, height=page.rect.height)
        fac = page.rect.width / pix.width
        mode = o.get("keep", "image")
        if mode == "image":
            newp.insert_image(newp.rect, pixmap=pix) if o.get("quality", "high") == "high" else newp.insert_image(newp.rect, stream=_jpeg(img, 60))
        else:
            newp.show_pdf_page(newp.rect, doc, pi)
        for text, (x0, y0, x1, y1) in lines:
            h = max(4.0, (y1 - y0) * fac)
            try:
                newp.insert_text((x0 * fac, y1 * fac - h * 0.2), text, fontsize=h * 0.9, render_mode=3, fontname="helv")
            except Exception:
                pass
        done += 1
    if not done:
        out.close(); doc.close()
        raise ConvertError("Every page already has text, so nothing was changed.")
    if o.get("title", True):
        out.set_metadata({**doc.metadata, "title": doc.metadata.get("title") or src.stem})
    out.save(str(dst), garbage=3, deflate=True)
    out.close(); doc.close()
    return f"Made {done} page(s) searchable" + (f", {empty} already had text" if empty else "")


def _jpeg(img, q):
    b = io.BytesIO(); img.save(b, "JPEG", quality=q); return b.getvalue()


def ocrpdf_start(paths, recursive, o, output):
    policy = OutputPolicy.from_dict(output)
    tasks = []
    for src, rel in expand_sources(paths, recursive):
        if src.suffix.lower() != ".pdf":
            continue

        def fn(ctx, task, src=src, rel=rel):
            dst = ctx.scratch_dir() / f"{src.stem} (searchable).pdf"
            note = ocr_pdf_one(ctx, src, dst, o)
            return Output(policy.place([dst], src, rel), note)
        tasks.append(Task(src.name, fn, src, heavy=True, rel=rel))
    if not tasks:
        raise UserError("Choose one or more PDF files.")
    return {"job": ops.start("ocrpdf", "Making PDFs searchable", tasks).id, "count": len(tasks)}


# ================================================================ background remover
def remove_bg(src: Path, dst: Path, o: dict) -> str:
    import cv2
    import numpy as np
    from PIL import Image, ImageOps
    im = ImageOps.exif_transpose(Image.open(src)).convert("RGB")
    maxside = _int(o.get("work"), 1400, 400, 4000)
    scale = min(1.0, maxside / max(im.size))
    small = im if scale == 1 else im.resize((max(1, round(im.width * scale)), max(1, round(im.height * scale))), Image.LANCZOS)
    a = np.asarray(small)
    h, w = a.shape[:2]
    method = o.get("method", "grabcut")
    tol = _int(o.get("tolerance"), 30, 1, 120)
    if method == "color":
        key = _rgb(o["color"]) if o.get("color") else None
        if key is None:
            edge = np.concatenate([a[0], a[-1], a[:, 0], a[:, -1]])
            key = tuple(int(x) for x in np.median(edge, axis=0))
        lab = cv2.cvtColor(a, cv2.COLOR_RGB2LAB).astype(np.float32)
        kl = cv2.cvtColor(np.uint8([[key]]), cv2.COLOR_RGB2LAB).astype(np.float32)[0, 0]
        dist = np.linalg.norm(lab - kl, axis=2)
        soft = max(1.0, tol * 0.5)
        alpha = np.clip((dist - tol) / soft, 0, 1)
        if o.get("connected", True):
            hard = (dist <= tol).astype(np.uint8)
            n, lbl = cv2.connectedComponents(hard)
            keep = np.zeros_like(hard)
            border = set(np.unique(np.concatenate([lbl[0], lbl[-1], lbl[:, 0], lbl[:, -1]]))) - {0}
            for b in border:
                keep[lbl == b] = 1
            alpha = np.where(keep == 1, alpha, 1.0)
        alpha = (alpha * 255).astype(np.uint8)
    else:
        mg = _int(o.get("margin"), 6, 1, 40)
        rect = (int(w * mg / 100), int(h * mg / 100), max(2, int(w * (100 - 2 * mg) / 100)), max(2, int(h * (100 - 2 * mg) / 100)))
        mask = np.zeros((h, w), np.uint8)
        bgd, fgd = np.zeros((1, 65), np.float64), np.zeros((1, 65), np.float64)
        cv2.grabCut(cv2.cvtColor(a, cv2.COLOR_RGB2BGR), mask, rect, bgd, fgd, _int(o.get("passes"), 5, 1, 12), cv2.GC_INIT_WITH_RECT)
        alpha = np.where((mask == 1) | (mask == 3), 255, 0).astype(np.uint8)
        n, lbl, stats, _ = cv2.connectedComponentsWithStats((alpha > 0).astype(np.uint8))
        if n > 1 and o.get("largest", True):
            big = 1 + int(np.argmax(stats[1:, cv2.CC_STAT_AREA]))
            alpha = np.where(lbl == big, 255, 0).astype(np.uint8)
        k = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (5, 5))
        alpha = cv2.morphologyEx(alpha, cv2.MORPH_CLOSE, k)
    grow = _int(o.get("grow"), 0, -20, 20)
    if grow:
        k = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (2 * abs(grow) + 1,) * 2)
        alpha = cv2.dilate(alpha, k) if grow > 0 else cv2.erode(alpha, k)
    if scale != 1:
        alpha = cv2.resize(alpha, im.size, interpolation=cv2.INTER_CUBIC)
    feather = float(o.get("feather", 1.5) or 0)
    if feather > 0:
        alpha = cv2.GaussianBlur(alpha, (0, 0), feather)
    rgba = np.dstack([np.asarray(im), alpha])
    res = Image.fromarray(rgba, "RGBA")
    if o.get("crop"):
        bb = res.getchannel("A").point(lambda v: 255 if v > 8 else 0).getbbox()
        if bb:
            pad = _int(o.get("pad"), 10, 0, 200)
            res = res.crop((max(0, bb[0] - pad), max(0, bb[1] - pad), min(res.width, bb[2] + pad), min(res.height, bb[3] + pad)))
    fill = o.get("fill", "none")
    if fill == "color":
        base = Image.new("RGBA", res.size, _rgb(o.get("fillcolor"), (255, 255, 255)) + (255,))
        res = Image.alpha_composite(base, res)
    ext = dst.suffix.lower()
    if ext in (".jpg", ".jpeg"):
        if fill != "color":
            base = Image.new("RGBA", res.size, (255, 255, 255, 255)); res = Image.alpha_composite(base, res)
        res.convert("RGB").save(dst, "JPEG", quality=95)
    else:
        res.save(dst, "WEBP" if ext == ".webp" else "PNG", **({"lossless": True} if ext == ".webp" else {}))
    return "Background removed" + (" (a plain colour background works best)" if method == "color" else "")


def bg_start(paths, recursive, o, output):
    policy = OutputPolicy.from_dict(output)
    ext = {"png": ".png", "webp": ".webp", "jpg": ".jpg"}.get(o.get("format", "png"), ".png")
    tasks = []
    for src, rel in expand_sources(paths, recursive):
        if src.suffix.lower() not in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif", ".heic", ".avif"):
            continue

        def fn(ctx, task, src=src, rel=rel):
            dst = ctx.scratch_dir() / f"{src.stem} (no background){ext}"
            note = remove_bg(src, dst, o)
            return Output(policy.place([dst], src, rel), note)
        tasks.append(Task(src.name, fn, src, rel=rel))
    if not tasks:
        raise UserError("Choose one or more pictures.")
    return {"job": ops.start("nobg", "Removing backgrounds", tasks).id, "count": len(tasks)}


# ================================================================ subtitles
TS = re.compile(r"(?:(\d+):)?(\d{1,2}):(\d{2})[.,](\d{1,3})")


def _t(m) -> float:
    return int(m.group(1) or 0) * 3600 + int(m.group(2)) * 60 + int(m.group(3)) + int(m.group(4).ljust(3, "0")) / 1000


def _fmt(sec: float, vtt: bool) -> str:
    sec = max(0.0, sec)
    ms = round(sec * 1000)
    h, r = divmod(ms, 3600000); m, r = divmod(r, 60000); s, ms = divmod(r, 1000)
    return f"{h:02d}:{m:02d}:{s:02d}{'.' if vtt else ','}{ms:03d}"


def read_subs(path: Path, charset: str = "auto") -> tuple[list[list], str]:
    raw = path.read_bytes()
    if charset == "auto":
        if raw.startswith(b"\xef\xbb\xbf"):
            enc = "utf-8-sig"
        elif raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
            enc = "utf-16"
        else:
            try:
                raw.decode("utf-8"); enc = "utf-8"
            except UnicodeDecodeError:
                from charset_normalizer import from_bytes
                best = from_bytes(raw).best()
                enc = best.encoding if best else "cp1252"
    else:
        enc = charset
    text = raw.decode(enc, "replace").replace("\r\n", "\n").replace("\r", "\n")
    cues: list[list] = []
    if path.suffix.lower() in (".ass", ".ssa"):
        for ln in text.split("\n"):
            if ln.startswith("Dialogue:"):
                p = ln.split(",", 9)
                if len(p) == 10:
                    mm = re.fullmatch(r"(\d+):(\d{2}):(\d{2})\.(\d{2})", p[1].strip()); m2 = re.fullmatch(r"(\d+):(\d{2}):(\d{2})\.(\d{2})", p[2].strip())
                    if mm and m2:
                        f = lambda m: int(m.group(1)) * 3600 + int(m.group(2)) * 60 + int(m.group(3)) + int(m.group(4)) / 100
                        cues.append([f(mm), f(m2), re.sub(r"\{[^}]*\}", "", p[9]).replace("\\N", "\n").replace("\\n", "\n")])
        return cues, enc
    if path.suffix.lower() == ".lrc":
        stamps = [(int(m.group(1)) * 60 + float(m.group(2)), m.group(3).strip()) for m in (re.match(r"^\[(\d+):(\d+(?:\.\d+)?)\]\s*(.*)$", ln.strip()) for ln in text.split("\n")) if m]
        for i, (t0, tx) in enumerate(stamps):
            if tx:
                cues.append([t0, stamps[i + 1][0] if i + 1 < len(stamps) else t0 + 4, tx])
        if cues:
            return cues, enc
    if path.suffix.lower() == ".sbv":
        text = re.sub(r"(\d+:\d{2}:\d{2}\.\d{3}),(\d+:\d{2}:\d{2}\.\d{3})", r"\1 --> \2", text)
    for blk in re.split(r"\n\s*\n", text):
        lines = blk.strip("\n").split("\n")
        for i, ln in enumerate(lines):
            if "-->" in ln:
                a, b = ln.split("-->", 1)
                ma, mb = TS.search(a), TS.search(b)
                if ma and mb:
                    cues.append([_t(ma), _t(mb), "\n".join(lines[i + 1:]).strip()])
                break
    if not cues:
        raise ConvertError("Fathom could not find any subtitles in this file.")
    return cues, enc


def fix_subs(cues: list[list], o: dict) -> tuple[list[list], list[str]]:
    notes = []
    shift = float(o.get("shift") or 0)
    if o.get("fps_from") and o.get("fps_to"):
        f = float(o["fps_from"]) / float(o["fps_to"])
        cues = [[a * f, b * f, t] for a, b, t in cues]; notes.append(f"frame rate {o['fps_from']}→{o['fps_to']}")
    if o.get("sync1") not in (None, "") and o.get("sync2") not in (None, "") and len(cues) > 1:
        # two points: first cue should start at sync1, last cue at sync2
        a0, a1 = cues[0][0], cues[-1][0]
        s1, s2 = float(o["sync1"]), float(o["sync2"])
        if a1 != a0:
            k = (s2 - s1) / (a1 - a0)
            cues = [[s1 + (a - a0) * k, s1 + (b - a0) * k, t] for a, b, t in cues]; notes.append("stretched to fit the two sync points")
    if shift:
        cues = [[a + shift, b + shift, t] for a, b, t in cues]; notes.append(f"shifted {shift:+g}s")
    out = []
    removed = 0
    for a, b, t in cues:
        if o.get("html", True):
            t = re.sub(r"</?[a-zA-Z][^>]*>", "", t)
        if o.get("sdh"):
            t = re.sub(r"\[[^\]]*\]|\([^)]*\)|^[A-Z][A-Z .'-]{1,20}:\s*", "", t, flags=re.M)
            t = re.sub(r"[♪♫#]+[^\n]*", "", t)
        if o.get("ads"):
            if re.search(r"opensubtitles|subscene|addic7ed|yify|www\.|downloaded from|sync(ed)? (and )?corrected", t, re.I):
                removed += 1; continue
        t = "\n".join(l.strip() for l in t.split("\n") if l.strip())
        if o.get("upper") == "sentence":
            t = t.capitalize() if t.isupper() else t
        if not t:
            removed += 1; continue
        out.append([a, b, t])
    if removed:
        notes.append(f"removed {removed} empty or unwanted cues")
    if o.get("dedupe", True):
        d = []
        for c in out:
            if d and d[-1][2] == c[2] and c[0] - d[-1][1] < 0.3:
                d[-1][1] = max(d[-1][1], c[1])
            else:
                d.append(c)
        if len(d) != len(out):
            notes.append(f"merged {len(out) - len(d)} duplicate cues")
        out = d
    out.sort(key=lambda c: c[0])
    if o.get("overlap", True):
        for i in range(len(out) - 1):
            if out[i][1] > out[i + 1][0]:
                out[i][1] = max(out[i][0] + 0.2, out[i + 1][0] - 0.02)
    mindur, maxdur = float(o.get("mindur") or 0), float(o.get("maxdur") or 0)
    for i, c in enumerate(out):
        if mindur and c[1] - c[0] < mindur:
            nxt = out[i + 1][0] - 0.02 if i + 1 < len(out) else c[0] + mindur
            c[1] = min(c[0] + mindur, nxt)
        if maxdur and c[1] - c[0] > maxdur:
            c[1] = c[0] + maxdur
    wrap = int(float(o.get("wrap") or 0))
    if wrap:
        import textwrap
        for c in out:
            c[2] = "\n".join(textwrap.wrap(c[2].replace("\n", " "), wrap)[:3])
    return out, notes


SUB_FORMATS = ["srt", "vtt", "ass", "ssa", "sbv", "lrc", "txt", "csv", "tsv", "json"]
SUB_IN = (".srt", ".vtt", ".ass", ".ssa", ".sbv", ".lrc")


def _ass_time(sec: float) -> str:
    cs = round(max(0.0, sec) * 100)
    h, r = divmod(cs, 360000); m, r = divmod(r, 6000); s_, c = divmod(r, 100)
    return f"{h}:{m:02d}:{s_:02d}.{c:02d}"


def write_subs(cues, path: Path, fmt: str, enc: str, style: dict | None = None):
    vtt = fmt == "vtt"
    st = style or {}
    if fmt == "txt":
        data = "\n".join(c[2].replace("\n", " ") for c in cues) + "\n"
    elif fmt in ("csv", "tsv"):
        sep = "," if fmt == "csv" else "\t"
        q = (lambda x: '"' + x.replace('"', '""') + '"') if fmt == "csv" else (lambda x: x.replace("\t", " "))
        data = "\n".join([f"start{sep}end{sep}text"] + [f"{a:.3f}{sep}{b:.3f}{sep}{q(t.replace(chr(10), ' '))}" for a, b, t in cues]) + "\n"
    elif fmt == "json":
        import json
        data = json.dumps([{"start": round(a, 3), "end": round(b, 3), "text": t} for a, b, t in cues], ensure_ascii=False, indent=2)
    elif fmt == "lrc":
        data = "\n".join(f"[{int(a // 60):02d}:{a % 60:05.2f}]{t.replace(chr(10), ' ')}" for a, b, t in cues) + "\n"
    elif fmt == "sbv":
        def t5(x):
            ms = round(max(0, x) * 1000); h, r = divmod(ms, 3600000); m, r = divmod(r, 60000); s_, ms = divmod(r, 1000)
            return f"{h}:{m:02d}:{s_:02d}.{ms:03d}"
        data = "\n".join(f"{t5(a)},{t5(b)}\n{t}\n" for a, b, t in cues)
    elif fmt in ("ass", "ssa"):
        font, size = st.get("font", "Arial"), st.get("size", 24)
        prim = "&H00" + _bgr(st.get("color", "#ffffff")); out = "&H00" + _bgr(st.get("outline", "#000000"))
        head = ("[Script Info]\nScriptType: v4.00+\nPlayResX: 1280\nPlayResY: 720\nWrapStyle: 0\n\n[V4+ Styles]\n"
                "Format: Name,Fontname,Fontsize,PrimaryColour,SecondaryColour,OutlineColour,BackColour,Bold,Italic,Underline,StrikeOut,ScaleX,ScaleY,Spacing,Angle,BorderStyle,Outline,Shadow,Alignment,MarginL,MarginR,MarginV,Encoding\n"
                f"Style: Default,{font},{size},{prim},&H000000FF,{out},&H80000000,0,0,0,0,100,100,0,0,1,2,1,2,40,40,30,1\n\n[Events]\nFormat: Layer,Start,End,Style,Name,MarginL,MarginR,MarginV,Effect,Text\n")
        data = head + "\n".join(f"Dialogue: 0,{_ass_time(a)},{_ass_time(b)},Default,,0,0,0,,{t.replace(chr(10), chr(92) + 'N')}" for a, b, t in cues) + "\n"
    else:
        parts = ["WEBVTT\n"] if vtt else []
        for i, (a, b, t) in enumerate(cues, 1):
            parts.append((f"{i}\n" if not vtt else "") + f"{_fmt(a, vtt)} --> {_fmt(b, vtt)}\n{t}\n")
        data = "\n".join(parts)
    if enc == "utf-8-sig":
        path.write_bytes(b"\xef\xbb\xbf" + data.encode("utf-8"))
    else:
        path.write_text(data, encoding=enc, errors="replace")


def subs_start(paths, recursive, o, output):
    policy = OutputPolicy.from_dict(output)
    fmt = o.get("format", "srt")
    tasks = []
    for src, rel in expand_sources(paths, recursive):
        if src.suffix.lower() not in SUB_IN:
            continue

        def fn(ctx, task, src=src, rel=rel):
            cues, enc = read_subs(src, o.get("charset", "auto"))
            cues, notes = fix_subs(cues, o)
            out_enc = {"keep": enc if enc != "utf-8-sig" else "utf-8", "utf8": "utf-8", "bom": "utf-8-sig", "latin1": "cp1252"}.get(o.get("encoding", "utf8"), "utf-8")
            dst = ctx.scratch_dir() / f"{src.stem} (fixed).{fmt}"
            write_subs(cues, dst, fmt, out_enc, {"font": o.get("font") or "Arial", "size": _int(o.get("fsize"), 24, 8, 120), "color": o.get("fcolor", "#ffffff"), "outline": o.get("foutline", "#000000")})
            return Output(policy.place([dst], src, rel), f"{len(cues)} cues; read as {enc}" + ("; " + ", ".join(notes) if notes else ""))
        tasks.append(Task(src.name, fn, src, rel=rel))
    if not tasks:
        raise UserError("Choose subtitle files (.srt, .vtt, .ass, .ssa, .sbv or .lrc).")
    return {"job": ops.start("subs", "Fixing subtitles", tasks).id, "count": len(tasks)}


def has_libass() -> bool:
    try:
        r = subprocess.run([ffmpeg_exe(), "-hide_banner", "-filters"], capture_output=True, text=True, timeout=20)
        return " subtitles " in r.stdout
    except Exception:
        return False


def burn_start(paths, recursive, o, output):
    """Burn a subtitle file into a video. The subtitle file is the one named in o['subs'], or a file with the same name beside the video."""
    policy = OutputPolicy.from_dict(output)
    if not has_libass():
        raise UserError("This copy of the video engine cannot draw subtitles on video. Use 'Add subtitles as a track' instead.")
    tasks = []
    for src, rel in expand_sources(paths, recursive):
        if src.suffix.lower() in SUB_IN + (".txt",):
            continue
        sub = Path(o["subs"]) if o.get("subs") else next((src.with_suffix(e) for e in SUB_IN if src.with_suffix(e).exists()), None)
        if not sub or not Path(sub).exists():
            continue

        def fn(ctx, task, src=src, rel=rel, sub=Path(sub)):
            from app.convert.eng_media import probe, run_ffmpeg
            d = ctx.scratch_dir()
            local = d / ("s" + sub.suffix.lower())
            shutil.copy(sub, local)
            if local.suffix == ".vtt":
                cues, _ = read_subs(local); local = d / "s.srt"; write_subs(cues, local, "srt", "utf-8")
            style = f"FontSize={_int(o.get('size'), 22, 8, 80)},PrimaryColour=&H{_bgr(o.get('color', '#ffffff'))},OutlineColour=&H{_bgr(o.get('outline', '#000000'))},Outline={_int(o.get('border'), 2, 0, 8)},Shadow=1,Alignment={ {'bottom': 2, 'top': 6, 'middle': 10}.get(o.get('pos', 'bottom'), 2) },MarginV={_int(o.get('margin'), 20, 0, 200)}"
            lp = str(local).replace("\\", "/").replace(":", "\\:").replace("'", "\\'")
            vf = f"subtitles=filename='{lp}':force_style='{style}'" if local.suffix != ".ass" else f"subtitles=filename='{lp}'"
            info = probe(ctx, src)
            dst = d / f"{src.stem} (subtitled).mp4"
            args = ["-i", src, "-vf", vf, "-c:v", "libx264", "-crf", str(_int(o.get("crf"), 20, 14, 32)), "-preset", "veryfast", "-pix_fmt", "yuv420p", "-c:a", "copy" if info["audio"] else "-an", "-movflags", "+faststart", dst]
            run_ffmpeg(ctx, args, info["duration"], "subtitles")
            return Output(policy.place([dst], src, rel), "Subtitles burned in")
        tasks.append(Task(src.name, fn, src, heavy=True, rel=rel))
    if not tasks:
        raise UserError("Choose a video, and a subtitle file (or keep one with the same name beside the video).")
    return {"job": ops.start("burn", "Burning subtitles", tasks).id, "count": len(tasks)}


def _bgr(c: str) -> str:
    r, g, b = _rgb(c, (255, 255, 255))
    return f"{b:02X}{g:02X}{r:02X}"


def subs_check() -> dict:
    return {"burn": has_libass()}


# ================================================================ translate subtitles (with the assistant's own model)
LANGS = {"en": "English", "zh": "Chinese", "de": "German", "es": "Spanish", "ru": "Russian", "ko": "Korean", "fr": "French", "ja": "Japanese", "pt": "Portuguese", "tr": "Turkish", "pl": "Polish", "ca": "Catalan", "nl": "Dutch", "ar": "Arabic", "sv": "Swedish", "it": "Italian", "id": "Indonesian", "hi": "Hindi", "fi": "Finnish", "vi": "Vietnamese", "he": "Hebrew", "uk": "Ukrainian", "el": "Greek", "ms": "Malay", "cs": "Czech", "ro": "Romanian", "da": "Danish", "hu": "Hungarian", "ta": "Tamil", "no": "Norwegian", "th": "Thai", "ur": "Urdu", "hr": "Croatian", "bg": "Bulgarian", "lt": "Lithuanian", "la": "Latin", "mi": "Maori", "ml": "Malayalam", "cy": "Welsh", "sk": "Slovak", "te": "Telugu", "fa": "Persian", "lv": "Latvian", "bn": "Bengali", "sr": "Serbian", "az": "Azerbaijani", "sl": "Slovenian", "kn": "Kannada", "et": "Estonian", "mk": "Macedonian", "br": "Breton", "eu": "Basque", "is": "Icelandic", "hy": "Armenian", "ne": "Nepali", "mn": "Mongolian", "bs": "Bosnian", "kk": "Kazakh", "sq": "Albanian", "sw": "Swahili", "gl": "Galician", "mr": "Marathi", "pa": "Punjabi", "si": "Sinhala", "km": "Khmer", "sn": "Shona", "yo": "Yoruba", "so": "Somali", "af": "Afrikaans", "oc": "Occitan", "ka": "Georgian", "be": "Belarusian", "tg": "Tajik", "sd": "Sindhi", "gu": "Gujarati", "am": "Amharic", "yi": "Yiddish", "lo": "Lao", "uz": "Uzbek", "fo": "Faroese", "ht": "Haitian Creole", "ps": "Pashto", "tk": "Turkmen", "nn": "Nynorsk", "mt": "Maltese", "sa": "Sanskrit", "lb": "Luxembourgish", "my": "Burmese", "bo": "Tibetan", "tl": "Tagalog", "mg": "Malagasy", "as": "Assamese", "tt": "Tatar", "haw": "Hawaiian", "ln": "Lingala", "ha": "Hausa", "ba": "Bashkir", "jw": "Javanese", "su": "Sundanese", "yue": "Cantonese", "zt": "Chinese (Traditional)", "ga": "Irish", "gd": "Scottish Gaelic", "ku": "Kurdish", "ky": "Kyrgyz", "ug": "Uyghur", "xh": "Xhosa", "zu": "Zulu", "ig": "Igbo", "rw": "Kinyarwanda", "ny": "Chichewa", "st": "Sesotho", "tn": "Tswana", "ti": "Tigrinya", "or": "Odia", "ckb": "Kurdish (Sorani)", "eo": "Esperanto", "fy": "Frisian", "co": "Corsican", "hmn": "Hmong", "ceb": "Cebuano", "ay": "Aymara", "qu": "Quechua", "gn": "Guarani", "dv": "Dhivehi", "sm": "Samoan", "to": "Tongan", "fj": "Fijian", "wo": "Wolof", "ak": "Akan", "ee": "Ewe", "lg": "Luganda", "om": "Oromo", "ff": "Fula", "tpi": "Tok Pisin", "bm": "Bambara", "ks": "Kashmiri", "doi": "Dogri", "kok": "Konkani", "mai": "Maithili", "mni": "Meiteilon (Manipuri)", "sat": "Santali", "bho": "Bhojpuri", "mwr": "Marwari", "rom": "Romani"}
_LINE = re.compile(r"^\s*(\d+)\s*[|:.)\]-]\s*(.*)$")


def _ask(system: str, user: str, cancel) -> str:
    from app.ai import llm
    try:
        return llm.complete(system, [{"role": "user", "content": user}], cancel, 4000)
    except llm.LLMError as e:
        raise ConvertError(str(e))


def translate_cues(cues: list[list], target: str, source: str, notes: str, batch: int, cancel, ask=None, progress=None, tone: str = "", glossary: str = "") -> tuple[list[str], int]:
    """Returns (translated texts in the same order, number of lines that had to stay in the original language)."""
    ask = ask or _ask
    lang = LANGS.get(target, target)
    src = f"from {LANGS.get(source, source)} " if source and source != "auto" else ""
    system = (f"You translate movie subtitles {src}into {lang}. Rules: keep the meaning, tone and register of speech; keep it short enough to read on screen; "
              "keep names; never add explanations. Each input line is 'number|text' where <br> marks a line break. "
              "Answer with exactly the same numbers, one 'number|translation' per line, keeping <br>, and nothing else." + (f" Style: {tone}." if tone else "") + (f" Always translate these terms exactly like this (original = translation): {glossary}." if glossary else "") + (f" Extra notes from the user: {notes}" if notes else ""))
    out = [""] * len(cues)
    kept = 0
    for start in range(0, len(cues), batch):
        if cancel is not None and cancel.is_set():
            raise ConvertError("Stopped.")
        chunk = list(range(start, min(len(cues), start + batch)))
        want = {i + 1: cues[i][2].replace("\n", "<br>") for i in chunk}
        got: dict[int, str] = {}
        for attempt in range(3):
            todo = [n for n in want if n not in got]
            if not todo:
                break
            ctx_lines = ""
            if start and attempt == 0 and out[start - 1]:
                ctx_lines = f"(For context only, do not translate: the line before was: {out[start - 1]})\n"
            reply = ask(system, ctx_lines + "\n".join(f"{n}|{want[n]}" for n in todo), cancel)
            for ln in reply.splitlines():
                m = _LINE.match(ln)
                if m and int(m.group(1)) in want and m.group(2).strip() and int(m.group(1)) not in got:
                    got[int(m.group(1))] = m.group(2).strip()
        for n in want:
            if n in got:
                out[n - 1] = got[n].replace("<br>", "\n").replace("<BR>", "\n")
            else:
                out[n - 1] = cues[n - 1][2]
                kept += 1
        if progress:
            progress(min(1.0, (start + len(chunk)) / len(cues)))
    return out, kept


def translate_start(paths, recursive, o, output):
    policy = OutputPolicy.from_dict(output)
    target = (o.get("target") or "").strip()
    if o.get("custom"):
        target = o["custom"].strip()
    if not target:
        raise UserError("Choose the language to translate into.")
    fmt = o.get("format", "srt")
    tasks = []
    for src, rel in expand_sources(paths, recursive):
        if src.suffix.lower() not in SUB_IN:
            continue

        def fn(ctx, task, src=src, rel=rel):
            cues, enc = read_subs(src, o.get("charset", "auto"))
            cues, _n = fix_subs(cues, {"html": True, "dedupe": False, "overlap": False})
            texts, kept = translate_cues(cues, target, o.get("source", "auto"), (o.get("notes") or "")[:500], _int(o.get("batch"), 20, 5, 60), ctx.cancel, progress=ctx.progress, tone=o.get("tone", ""), glossary="; ".join(l.strip() for l in (o.get("glossary") or "").splitlines() if "=" in l)[:1500])
            mode = o.get("layout", "only")
            result = []
            for (a, b, orig), tr in zip(cues, texts):
                result.append([a, b, tr if mode == "only" else f"{tr}\n{orig}" if mode == "top" else f"{orig}\n{tr}"])
            if _int(o.get("wrap"), 0, 0, 100) or o.get("tidy"):
                result, _n = fix_subs(result, {"html": False, "dedupe": bool(o.get("tidy")), "overlap": bool(o.get("tidy")), "wrap": o.get("wrap"), "mindur": o.get("mindur"), "shift": o.get("shift")})
            code = target if target in LANGS else re.sub(r"\W+", "", target)[:12].lower()
            dst = ctx.scratch_dir() / f"{src.stem}.{code}.{fmt}"
            write_subs(result, dst, fmt, "utf-8-sig" if o.get("bom") else "utf-8")
            return Output(policy.place([dst], src, rel), f"{len(result)} lines into {LANGS.get(target, target)}" + (f"; {kept} lines could not be translated and were left as they were" if kept else ""))
        tasks.append(Task(src.name, fn, src, heavy=True, rel=rel))
    if not tasks:
        raise UserError("Choose subtitle files (.srt, .vtt, .ass, .ssa, .sbv or .lrc).")
    return {"job": ops.start("translate", "Translating subtitles", tasks).id, "count": len(tasks)}


# ================================================================ translate documents (text, Word, PowerPoint, Excel, web pages, CSV, JSON, PDF)
DOC_EXT = (".txt", ".md", ".html", ".htm", ".csv", ".json", ".docx", ".pptx", ".xlsx", ".pdf")


def _tr_list(texts: list[str], target: str, source: str, o: dict, ctx) -> tuple[list[str], int]:
    idx = [i for i, t in enumerate(texts) if re.search(r"\w", t or "")]
    if len(idx) > 6000:
        raise ConvertError("This file has too many pieces of text to translate in one go (more than 6000). Split it first.")
    cues = [[0, 0, texts[i]] for i in idx]
    tr, kept = translate_cues(cues, target, source, (o.get("notes") or "")[:500], _int(o.get("batch"), 15, 3, 40), ctx.cancel, progress=ctx.progress, tone=o.get("tone", ""), glossary="; ".join(l.strip() for l in (o.get("glossary") or "").splitlines() if "=" in l)[:1500])
    out = list(texts)
    for i, t in zip(idx, tr):
        out[i] = t
    return out, kept


def translate_doc(src: Path, dst_dir: Path, target: str, o: dict, ctx) -> tuple[Path, int, int]:
    ext = src.suffix.lower()
    source = o.get("source", "auto")
    code = target if target in LANGS else re.sub(r"\W+", "", target)[:12].lower()
    dst = dst_dir / f"{src.stem}.{code}{ext if ext != '.pdf' else '.txt'}"
    both = o.get("layout", "only") != "only"
    kept = 0
    if ext in (".txt", ".md", ".pdf"):
        if ext == ".pdf":
            try:
                import pymupdf as fitz
            except ImportError:
                import fitz
            doc = fitz.open(str(src))
            text = "\n\n".join(p.get_text() for p in doc)
            if len(text.strip()) < 20:
                raise ConvertError("This PDF has no text to translate (it is a scan). Use Searchable PDF first.")
        else:
            text = src.read_bytes().decode("utf-8", "replace")
        paras = re.split(r"\n\s*\n", text.strip())
        res, kept = _tr_list([" ".join(p.split()) if ext == ".pdf" else p for p in paras], target, source, o, ctx)
        dst.write_text("\n\n".join((f"{a}\n{b}" if both and a != b else a) for a, b in zip(res, paras)) + "\n", encoding="utf-8")
        return dst, len(paras), kept
    if ext in (".html", ".htm"):
        from html.parser import HTMLParser
        raw = src.read_bytes().decode("utf-8", "replace")
        pieces: list[tuple[bool, str]] = []

        class P(HTMLParser):
            def __init__(self):
                super().__init__(convert_charrefs=True); self.skip = 0
            def handle_starttag(self, tag, attrs):
                if tag in ("script", "style", "code", "pre"): self.skip += 1
                pieces.append((False, self.get_starttag_text()))
            def handle_endtag(self, tag):
                if tag in ("script", "style", "code", "pre"): self.skip = max(0, self.skip - 1)
                pieces.append((False, f"</{tag}>"))
            def handle_startendtag(self, tag, attrs): pieces.append((False, self.get_starttag_text()))
            def handle_data(self, d): pieces.append((not self.skip and bool(d.strip()), d))
            def handle_comment(self, d): pieces.append((False, f"<!--{d}-->"))
            def handle_decl(self, d): pieces.append((False, f"<!{d}>"))
        P().feed(raw)
        idx = [i for i, (t, _s) in enumerate(pieces) if t]
        keep = [(re.match(r"\s*", pieces[i][1]).group(0), re.search(r"\s*$", pieces[i][1]).group(0)) for i in idx]
        res, kept = _tr_list([pieces[i][1].strip() for i in idx], target, source, o, ctx)
        import html as _h
        for i, r, (lead, trail) in zip(idx, res, keep):
            pieces[i] = (False, lead + _h.escape(r, quote=False) + trail)
        dst.write_text("".join(s for _t, s in pieces), encoding="utf-8")
        return dst, len(idx), kept
    if ext == ".csv":
        import csv
        text = src.read_bytes().decode("utf-8-sig", "replace")
        rows = list(csv.reader(io.StringIO(text), dialect=csv.Sniffer().sniff(text[:4000]) if text.strip() else csv.excel))
        flat = [c for r in rows for c in r]
        res, kept = _tr_list([c if re.search(r"[A-Za-z\u0080-￿]{2,}", c) else "" for c in flat], target, source, o, ctx)
        it = iter(res); fl = iter(flat); out = []
        for r in rows:
            nr = []
            for _c in r:
                a, b = next(it), next(fl)
                nr.append(a if a else b)
            out.append(nr)
        buf = io.StringIO(); csv.writer(buf, lineterminator="\n").writerows(out)
        dst.write_text(buf.getvalue(), encoding="utf-8")
        return dst, len(flat), kept
    if ext == ".json":
        import json
        data = json.loads(src.read_bytes().decode("utf-8-sig", "replace"))
        leaves: list = []

        def walk(x):
            if isinstance(x, dict):
                for k in x:
                    if isinstance(x[k], str): leaves.append((x, k))
                    else: walk(x[k])
            elif isinstance(x, list):
                for i, v in enumerate(x):
                    if isinstance(v, str): leaves.append((x, i))
                    else: walk(v)
        walk(data)
        res, kept = _tr_list([c[k] for c, k in leaves], target, source, o, ctx)
        for (c, k), r in zip(leaves, res):
            c[k] = r
        dst.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        return dst, len(leaves), kept
    if ext == ".docx":
        import docx
        d = docx.Document(str(src))
        paras = list(d.paragraphs)
        for t in d.tables:
            for row in t.rows:
                for cell in row.cells:
                    paras += cell.paragraphs
        for sec in d.sections:
            paras += sec.header.paragraphs + sec.footer.paragraphs
        seen, uniq = set(), []
        for p in paras:
            if id(p._p) not in seen:
                seen.add(id(p._p)); uniq.append(p)
        res, kept = _tr_list([p.text for p in uniq], target, source, o, ctx)
        for p, r in zip(uniq, res):
            if r != p.text and p.runs:
                orig = p.text
                p.runs[0].text = f"{r}\n{orig}" if both else r
                for run in p.runs[1:]:
                    run.text = ""
        d.save(str(dst))
        return dst, len(uniq), kept
    if ext == ".pptx":
        from pptx import Presentation
        pr = Presentation(str(src))
        paras = []

        def frames(shapes):
            for sh in shapes:
                if sh.shape_type == 6 and hasattr(sh, "shapes"):
                    yield from frames(sh.shapes)
                if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                    paras.extend(sh.text_frame.paragraphs)
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        for c in row.cells:
                            paras.extend(c.text_frame.paragraphs)
        for sl in pr.slides:
            frames_list = list(frames(sl.shapes))
            if sl.has_notes_slide:
                paras.extend(sl.notes_slide.notes_text_frame.paragraphs)
        texts = ["".join(r.text for r in p.runs) for p in paras]
        res, kept = _tr_list(texts, target, source, o, ctx)
        for p, a, r in zip(paras, texts, res):
            if p.runs and r != a:
                p.runs[0].text = f"{r} / {a}" if both else r
                for run in p.runs[1:]:
                    run.text = ""
        pr.save(str(dst))
        return dst, len(paras), kept
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(src))
        cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and not c.value.startswith("=")]
        res, kept = _tr_list([c.value for c in cells], target, source, o, ctx)
        for c, r in zip(cells, res):
            c.value = f"{r}\n{c.value}" if both and r != c.value else r
        wb.save(str(dst))
        return dst, len(cells), kept
    raise ConvertError("This kind of file cannot be translated.")


def html_unescape(s: str) -> str:
    import html as _h
    return _h.unescape(s)


def translate_docs_start(paths, recursive, o, output):
    policy = OutputPolicy.from_dict(output)
    target = (o.get("custom") or o.get("target") or "").strip()
    if not target:
        raise UserError("Choose the language to translate into.")
    tasks = []
    for src, rel in expand_sources(paths, recursive):
        if src.suffix.lower() not in DOC_EXT:
            continue

        def fn(ctx, task, src=src, rel=rel):
            dst, n, kept = translate_doc(src, ctx.scratch_dir(), target, o, ctx)
            return Output(policy.place([dst], src, rel), f"{n} pieces of text into {LANGS.get(target, target)}" + (f"; {kept} could not be translated and were left as they were" if kept else "") + ("; mixed formatting inside a paragraph is simplified" if src.suffix.lower() in (".docx", ".pptx") else ""))
        tasks.append(Task(src.name, fn, src, heavy=True, rel=rel))
    if not tasks:
        raise UserError("Choose text, Word, PowerPoint, Excel, web page, CSV, JSON or PDF files.")
    return {"job": ops.start("translatedocs", "Translating documents", tasks).id, "count": len(tasks)}


# ================================================================ picture colours (palette), contact sheet, sound pictures
def _hsl(r, g, b):
    import colorsys
    h, l, s = colorsys.rgb_to_hls(r / 255, g / 255, b / 255)
    return round(h * 360), round(s * 100), round(l * 100)


def palette_one(src: Path, dst_dir: Path, o: dict) -> list[Path]:
    import cv2
    import numpy as np
    from PIL import Image, ImageDraw, ImageFont, ImageOps
    im = ImageOps.exif_transpose(Image.open(src)).convert("RGBA")
    im.thumbnail((256, 256))
    a = np.asarray(im).reshape(-1, 4)
    a = a[a[:, 3] > 128][:, :3].astype(np.float32)
    if o.get("skip_white", False):
        a = a[(a.min(axis=1) < 235)]
    if o.get("skip_black", False):
        a = a[(a.max(axis=1) > 25)]
    if len(a) < 10:
        raise ConvertError("Not enough colour in this picture.")
    k = _int(o.get("count"), 6, 2, 24)
    k = min(k, len(a))
    cv2.setRNGSeed(1)
    _c, labels, centers = cv2.kmeans(a, k, None, (cv2.TERM_CRITERIA_EPS + cv2.TERM_CRITERIA_MAX_ITER, 30, 0.5), 3, cv2.KMEANS_PP_CENTERS)
    counts = np.bincount(labels.flatten(), minlength=k)
    cols = [(tuple(int(round(v)) for v in centers[i]), counts[i] / counts.sum()) for i in range(k)]
    order = o.get("order", "share")
    if order == "hue":
        cols.sort(key=lambda c: _hsl(*c[0])[0])
    elif order == "light":
        cols.sort(key=lambda c: _hsl(*c[0])[2])
    else:
        cols.sort(key=lambda c: -c[1])
    outs = []
    fmts = o.get("formats") or ["png", "txt"]
    hexs = ["#%02x%02x%02x" % c for c, _ in cols]
    stem = src.stem
    if "png" in fmts:
        w, h = _int(o.get("sw"), 120, 40, 400), _int(o.get("sh"), 160, 40, 600)
        vertical = o.get("layout") == "vertical"
        sheet = Image.new("RGB", (h if vertical else w * len(cols), w * len(cols) if vertical else h), (255, 255, 255))
        d = ImageDraw.Draw(sheet)
        try:
            font = ImageFont.load_default(size=max(10, w // 9))
        except TypeError:
            font = ImageFont.load_default()
        for i, (c, share) in enumerate(cols):
            box = (0, i * w, h, (i + 1) * w) if vertical else (i * w, 0, (i + 1) * w, h)
            d.rectangle(box, fill=c)
            if o.get("label", True):
                tc = (0, 0, 0) if _hsl(*c)[2] > 55 else (255, 255, 255)
                d.text((box[0] + 8, box[3] - max(24, w // 5)), hexs[i], fill=tc, font=font)
        p = dst_dir / f"{stem} (palette).png"; sheet.save(p); outs.append(p)
    if "txt" in fmts:
        p = dst_dir / f"{stem} (palette).txt"
        p.write_text("\n".join(f"{hexs[i]}  rgb({c[0]}, {c[1]}, {c[2]})  hsl({_hsl(*c)[0]}, {_hsl(*c)[1]}%, {_hsl(*c)[2]}%)  {sh * 100:.1f}%" for i, (c, sh) in enumerate(cols)) + "\n", encoding="utf-8"); outs.append(p)
    if "css" in fmts:
        p = dst_dir / f"{stem} (palette).css"
        p.write_text(":root {\n" + "\n".join(f"  --color-{i + 1}: {hexs[i]};" for i in range(len(cols))) + "\n}\n", encoding="utf-8"); outs.append(p)
    if "json" in fmts:
        import json
        p = dst_dir / f"{stem} (palette).json"
        p.write_text(json.dumps([{"hex": hexs[i], "rgb": list(c), "hsl": list(_hsl(*c)), "share": round(sh, 4)} for i, (c, sh) in enumerate(cols)], indent=2), encoding="utf-8"); outs.append(p)
    if "gpl" in fmts:
        p = dst_dir / f"{stem} (palette).gpl"
        p.write_text("GIMP Palette\nName: " + stem + "\n#\n" + "\n".join(f"{c[0]:3d} {c[1]:3d} {c[2]:3d} {hexs[i]}" for i, (c, _) in enumerate(cols)) + "\n", encoding="utf-8"); outs.append(p)
    return outs


def palette_start(paths, recursive, o, output):
    policy = OutputPolicy.from_dict(output)
    tasks = []
    for src, rel in expand_sources(paths, recursive):
        if src.suffix.lower() not in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif", ".heic", ".avif"):
            continue

        def fn(ctx, task, src=src, rel=rel):
            outs = palette_one(src, ctx.scratch_dir(), o)
            return Output(policy.place(outs, src, rel), "")
        tasks.append(Task(src.name, fn, src, rel=rel))
    if not tasks:
        raise UserError("Choose one or more pictures.")
    return {"job": ops.start("palette", "Finding colours", tasks).id, "count": len(tasks)}


PAGES = {"a4": (2480, 3508), "letter": (2550, 3300), "a3": (3508, 4961), "a5": (1748, 2480)}


def contact_start(paths, recursive, o, output):
    policy = OutputPolicy.from_dict(output)
    files = [(s, r) for s, r in expand_sources(paths, recursive) if s.suffix.lower() in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif", ".heic", ".avif")]
    if not files:
        raise UserError("Choose some pictures first.")
    first = files[0][0]

    def fn(ctx, task):
        from PIL import Image, ImageDraw, ImageFont, ImageOps
        cols, th, gap, pad = _int(o.get("cols"), 5, 1, 20), _int(o.get("thumb"), 300, 60, 1200), _int(o.get("gap"), 16, 0, 100), _int(o.get("pad"), 40, 0, 300)
        bg, fg = _rgb(o.get("bg"), (255, 255, 255)), _rgb(o.get("fg"), (30, 30, 30))
        labels, title = bool(o.get("labels", True)), (o.get("title") or "").strip()
        fit = o.get("fit", "fit")
        try:
            font = ImageFont.load_default(size=max(11, th // 18))
        except TypeError:
            font = ImageFont.load_default()
        lab_h = max(11, th // 18) + 10 if labels else 0
        cell_w, cell_h = th, th + lab_h
        title_h = (max(28, th // 8) + 20) if title else 0
        per = o.get("page", "single")
        if per == "single":
            rows_per = None
        else:
            pw, ph = PAGES.get(per, PAGES["a4"])
            if o.get("landscape"):
                pw, ph = ph, pw
            cols = max(1, min(cols, (pw - 2 * pad + gap) // (cell_w + gap) or 1))
            rows_per = max(1, (ph - 2 * pad - title_h + gap) // (cell_h + gap))
        chunks = [files] if rows_per is None else [files[i:i + cols * rows_per] for i in range(0, len(files), cols * rows_per)]
        sheets = []
        try:
            tfont = ImageFont.load_default(size=max(28, th // 8))
        except TypeError:
            tfont = ImageFont.load_default()
        for ci, chunk in enumerate(chunks):
            rows = -(-len(chunk) // cols)
            if rows_per is None:
                W, H = 2 * pad + cols * cell_w + (cols - 1) * gap, 2 * pad + title_h + rows * cell_h + (rows - 1) * gap
            else:
                W, H = pw, ph
            sh = Image.new("RGB", (W, H), bg)
            d = ImageDraw.Draw(sh)
            if title:
                d.text((pad, pad // 2 + 4), title + (f"  ({ci + 1}/{len(chunks)})" if len(chunks) > 1 else ""), fill=fg, font=tfont)
            for i, (src, _r) in enumerate(chunk):
                ctx.check()
                r, c = divmod(i, cols)
                x, y = pad + c * (cell_w + gap), pad + title_h + r * (cell_h + gap)
                try:
                    im = ImageOps.exif_transpose(Image.open(src)).convert("RGB")
                except Exception:
                    continue
                if fit == "crop":
                    im = ImageOps.fit(im, (th, th), Image.LANCZOS)
                else:
                    im.thumbnail((th, th), Image.LANCZOS)
                sh.paste(im, (x + (th - im.width) // 2, y + (th - im.height) // 2))
                if o.get("border"):
                    d.rectangle([x + (th - im.width) // 2 - 1, y + (th - im.height) // 2 - 1, x + (th + im.width) // 2, y + (th + im.height) // 2], outline=fg)
                if labels:
                    name = src.stem if len(src.stem) <= 28 else src.stem[:25] + "…"
                    tw = d.textlength(name, font=font)
                    d.text((x + (th - tw) / 2, y + th + 4), name, fill=fg, font=font)
            sheets.append(sh)
            ctx.progress((ci + 1) / len(chunks))
        fmt = o.get("format", "pdf" if per != "single" else "png")
        dst = ctx.scratch_dir() / f"contact sheet.{fmt if fmt != 'jpg' else 'jpg'}"
        if fmt == "pdf":
            sheets[0].save(dst, "PDF", save_all=True, append_images=sheets[1:], resolution=_int(o.get("dpi"), 300, 72, 600))
            return Output(policy.place([dst], first, Path(".")), f"{len(files)} pictures on {len(sheets)} page(s)")
        outs = []
        for i, sh in enumerate(sheets):
            p = ctx.scratch_dir() / (f"contact sheet {i + 1}.{fmt}" if len(sheets) > 1 else dst.name)
            sh.save(p, "JPEG" if fmt == "jpg" else fmt.upper(), **({"quality": 92} if fmt in ("jpg", "webp") else {}))
            outs.append(p)
        return Output(policy.place(outs, first, Path(".")), f"{len(files)} pictures on {len(sheets)} sheet(s)")
    return {"job": ops.start("contact", "Making a contact sheet", [Task("Contact sheet", fn, first, rel=Path("."))]).id, "count": len(files)}


def wave_start(paths, recursive, o, output):
    policy = OutputPolicy.from_dict(output)
    tasks = []
    for src, rel in expand_sources(paths, recursive):
        if src.suffix.lower() in (".txt", ".srt", ".pdf", ".png", ".jpg", ".zip"):
            continue

        def fn(ctx, task, src=src, rel=rel):
            from app.convert.eng_media import probe, run_ffmpeg
            info = probe(ctx, src)
            if not info.get("audio"):
                raise ConvertError("This file has no sound.")
            kind = o.get("kind", "wave")
            w, h = _int(o.get("w"), 1600, 200, 8000), _int(o.get("h"), 400, 60, 4000)
            col = "%02x%02x%02x" % _rgb(o.get("fg"), (0, 160, 255)); bgc = "%02x%02x%02x" % _rgb(o.get("bg"), (16, 18, 32))
            if kind == "spectrum":
                flt = f"showspectrumpic=s={w}x{h}:legend={1 if o.get('legend', True) else 0}:color={o.get('palette', 'intensity')}:scale={o.get('scale', 'log')}"
            else:
                flt = f"aformat=channel_layouts=mono,showwavespic=s={w}x{h}:colors=#{col}:scale={o.get('scale', 'lin')}:draw={o.get('draw', 'scale')}" if not o.get("split") else f"showwavespic=s={w}x{h}:split_channels=1:colors=#{col}|#{col}:scale={o.get('scale', 'lin')}:draw={o.get('draw', 'scale')}"
            dst = ctx.scratch_dir() / f"{src.stem} ({'spectrogram' if kind == 'spectrum' else 'waveform'}).png"
            raw = ctx.scratch_dir() / "raw.png"
            run_ffmpeg(ctx, ["-i", src, "-filter_complex", flt, "-frames:v", "1", raw], info["duration"], "waveform")
            from PIL import Image
            im = Image.open(raw).convert("RGBA")
            if kind != "spectrum" and not o.get("transparent"):
                base = Image.new("RGBA", im.size, tuple(int(bgc[i:i + 2], 16) for i in (0, 2, 4)) + (255,))
                im = Image.alpha_composite(base, im)
            im.save(dst)
            return Output(policy.place([dst], src, rel), "")
        tasks.append(Task(src.name, fn, src, heavy=True, rel=rel))
    if not tasks:
        raise UserError("Choose songs or videos.")
    return {"job": ops.start("wave", "Drawing sound pictures", tasks).id, "count": len(tasks)}


# ================================================================ back up and restore Fathom's own settings
def backup_settings(output: dict) -> dict:
    import json
    from datetime import datetime as _dt
    from app.paths import data_root
    d = Path((output or {}).get("dir") or (Path.home() / "Documents" / "Fathom backups"))
    d.mkdir(parents=True, exist_ok=True)
    pack = {"fathom_backup": 1, "made": _dt.now().isoformat(timespec="seconds")}
    for name in ("prefs.json", "watch.json"):
        f = data_root() / name
        if f.exists():
            try:
                pack[name] = json.loads(f.read_text(encoding="utf-8"))
            except Exception:
                pass
    dst = d / f"fathom-settings-{_dt.now():%Y-%m-%d_%H%M}.json"
    dst.write_text(json.dumps(pack, ensure_ascii=False, indent=1), encoding="utf-8")
    return {"path": str(dst), "items": [k for k in pack if k.endswith(".json")]}


def restore_settings(path: str) -> dict:
    import json, os
    from app.paths import data_root
    p = Path(path)
    if not p.is_file():
        raise UserError("Choose a Fathom settings backup file.")
    try:
        pack = json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        raise UserError("This is not a Fathom settings backup.")
    if not isinstance(pack, dict) or pack.get("fathom_backup") != 1:
        raise UserError("This is not a Fathom settings backup.")
    done = []
    for name in ("prefs.json", "watch.json"):
        if name in pack:
            f = data_root() / name
            if f.exists():
                shutil.copy2(f, f.with_suffix(".before-restore"))
            tmp = f.with_suffix(".tmp")
            tmp.write_text(json.dumps(pack[name], ensure_ascii=False), encoding="utf-8")
            os.replace(tmp, f)
            done.append(name)
    return {"restored": done}
